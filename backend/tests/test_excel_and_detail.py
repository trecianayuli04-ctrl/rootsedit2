"""Iteration 6 — Excel export (RBAC + content), pos/summary enrichment,
pos/close carries detail, settings close_report_detail toggle."""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

BASE = os.environ.get("REACT_APP_BACKEND_URL", "https://roots-edit.preview.emergentagent.com").rstrip("/")
API = f"{BASE}/api"
PWD = "Rollow@2026"
USERS = ["owner", "finance", "marketing", "crew", "manager", "spv"]


def login(email_user):
    email = f"{email_user}@rollow.id"
    r = requests.post(f"{API}/auth/login", json={"email": email, "password": PWD}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text}"
    return r.json()["access_token"]


@pytest.fixture(scope="module")
def tokens():
    return {u: login(u) for u in USERS}


def H(tok):
    return {"Authorization": f"Bearer {tok}"}


# ---------- RBAC for /reports/{rtype}/export ----------
@pytest.mark.parametrize("rtype", ["sales", "finance", "nota"])
@pytest.mark.parametrize("role", ["owner", "finance"])
def test_export_allowed_for_owner_finance(tokens, rtype, role):
    r = requests.get(f"{API}/reports/{rtype}/export?period=month", headers=H(tokens[role]), timeout=60)
    assert r.status_code == 200, f"{role}/{rtype}: {r.status_code} {r.text[:200]}"
    assert r.headers.get("content-type", "").startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    ), r.headers.get("content-type")
    assert len(r.content) > 500, "xlsx seems empty"


@pytest.mark.parametrize("rtype", ["sales", "finance", "nota"])
@pytest.mark.parametrize("role", ["marketing", "crew", "manager", "spv"])
def test_export_forbidden_for_other_roles(tokens, rtype, role):
    r = requests.get(f"{API}/reports/{rtype}/export?period=month", headers=H(tokens[role]), timeout=30)
    assert r.status_code == 403, f"{role}/{rtype} expected 403 got {r.status_code}"


def test_export_404_for_non_exportable_rtype(tokens):
    r = requests.get(f"{API}/reports/inventory/export?period=month", headers=H(tokens["owner"]), timeout=30)
    assert r.status_code == 404


# ---------- Seed sales via crew for content tests ----------
@pytest.fixture(scope="module")
def seeded_sales(tokens):
    crew = tokens["crew"]
    # Register
    cur = requests.get(f"{API}/register/current", headers=H(crew), timeout=30).json()
    if not cur or not cur.get("id"):
        o = requests.post(f"{API}/register/open", headers=H(crew), json={"opening_cash": 100000}, timeout=30)
        assert o.status_code == 200, o.text
    menus = requests.get(f"{API}/menus", headers=H(crew), timeout=30).json()
    assert menus, "no menus seeded"
    m = menus[0]
    created = []
    for method in ("Cash", "QRIS"):
        payload = {
            "items": [{"menu_id": m["id"], "name": m["name"], "price": m["price"], "qty": 1}],
            "payment_method": method,
            "channel": "kasir",
        }
        r = requests.post(f"{API}/sales", headers=H(crew), json=payload, timeout=30)
        assert r.status_code == 200, f"create sale {method}: {r.status_code} {r.text}"
        created.append(r.json())
    return {"menu": m, "sales": created}


# ---------- Content check: sales workbook ----------
def test_sales_export_content(tokens, seeded_sales):
    r = requests.get(f"{API}/reports/sales/export?period=month", headers=H(tokens["owner"]), timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Rekap Menu" in wb.sheetnames
    assert "Detail Transaksi" in wb.sheetnames
    ws = wb["Rekap Menu"]
    # Row 5 should be headers
    headers = [ws.cell(row=5, column=c).value for c in range(1, 10)]
    assert headers == ["Menu", "Kategori", "Sub Kategori", "Qty Terjual", "% Qty",
                       "Penjualan", "% Penjualan", "COGS", "Margin"], headers
    # find TOTAL row
    found_total = False
    for row in ws.iter_rows(min_row=6, max_col=1, values_only=True):
        if row[0] == "TOTAL":
            found_total = True
            break
    assert found_total, "TOTAL row missing in Rekap Menu"

    ws2 = wb["Detail Transaksi"]
    headers2 = [ws2.cell(row=5, column=c).value for c in range(1, 16)]
    for required in ["Metode Pembayaran", "Status", "Subtotal", "Diskon", "Pajak", "Total", "HPP", "Laba"]:
        assert required in headers2, f"{required} missing in Detail Transaksi headers: {headers2}"


def test_finance_export_sheets(tokens):
    r = requests.get(f"{API}/reports/finance/export?period=month", headers=H(tokens["owner"]), timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Ringkasan" in wb.sheetnames
    assert "Detail Transaksi" in wb.sheetnames


def test_nota_export_sheets(tokens):
    r = requests.get(f"{API}/reports/nota/export?period=month", headers=H(tokens["owner"]), timeout=60)
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert "Rekap per Nota" in wb.sheetnames


# ---------- pos/summary enriched ----------
def test_pos_summary_by_method_detail(tokens, seeded_sales):
    r = requests.get(f"{API}/pos/summary", headers=H(tokens["crew"]), timeout=30)
    assert r.status_code == 200
    data = r.json()
    for k in ("by_method_detail", "gross", "tax", "service", "net", "void_count", "void_amount"):
        assert k in data, f"missing {k}"
    bmd = data["by_method_detail"]
    assert "Cash" in bmd and "QRIS" in bmd, f"methods: {list(bmd.keys())}"
    for pm in ("Cash", "QRIS"):
        assert "count" in bmd[pm] and "amount" in bmd[pm]
        assert bmd[pm]["count"] >= 1
        assert bmd[pm]["amount"] > 0


# ---------- pos/close carries detail ----------
def test_pos_close_carries_detail(tokens):
    crew = tokens["crew"]
    summ = requests.get(f"{API}/pos/summary", headers=H(crew), timeout=30).json()
    kas_exp = summ["cash_ledger"]["kas_sistem"]
    r = requests.post(f"{API}/pos/close", headers=H(crew),
                      json={"kas_aktual": kas_exp, "selisih_note": "test"}, timeout=30)
    assert r.status_code == 200, r.text
    rec = r.json()
    for k in ("by_method_detail", "gross", "tax", "net", "void_count"):
        assert k in rec, f"close missing {k}: keys={list(rec.keys())}"


# ---------- Settings close_report_detail persist ----------
def test_settings_close_report_detail_persists(tokens):
    tok = tokens["owner"]
    s = requests.get(f"{API}/settings", headers=H(tok), timeout=30).json()
    assert "close_report_detail" in s
    orig = s["close_report_detail"]
    # Flip to False
    body = {**s, "close_report_detail": False}
    body.pop("_id", None)
    r = requests.put(f"{API}/settings", headers=H(tok), json=body, timeout=30)
    assert r.status_code == 200, r.text
    s2 = requests.get(f"{API}/settings", headers=H(tok), timeout=30).json()
    assert s2["close_report_detail"] is False
    # Restore
    body = {**s2, "close_report_detail": True}
    body.pop("_id", None)
    requests.put(f"{API}/settings", headers=H(tok), json=body, timeout=30)
    s3 = requests.get(f"{API}/settings", headers=H(tok), timeout=30).json()
    assert s3["close_report_detail"] is True
    assert orig is True  # default
