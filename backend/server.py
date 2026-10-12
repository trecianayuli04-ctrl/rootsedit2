from fastapi import FastAPI, APIRouter, Depends, HTTPException, Request
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials
from fastapi.responses import StreamingResponse
from dotenv import load_dotenv
from starlette.middleware.cors import CORSMiddleware
from motor.motor_asyncio import AsyncIOMotorClient
from pydantic import BaseModel, Field
from typing import List, Optional
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from pathlib import Path
import os, uuid, logging, bcrypt, jwt, httpx
import io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

ROOT_DIR = Path(__file__).parent
load_dotenv(ROOT_DIR / '.env')

client = AsyncIOMotorClient(os.environ['MONGO_URL'])
db = client[os.environ['DB_NAME']]
JWT_SECRET = os.environ['JWT_SECRET']
TZ = ZoneInfo("Asia/Jakarta")
OUTLET = "rollow-main"  # multi-outlet ready: every document carries outlet_id
LOCATIONS = ["gudang", "bar", "kitchen"]

app = FastAPI(title="RLW Outlet Management")
api = APIRouter(prefix="/api")
bearer = HTTPBearer(auto_error=False)
logger = logging.getLogger("rollow")

# ---------------- Roles / permissions ----------------
ALL = ["owner", "manager", "supervisor", "finance", "marketing", "crew"]
OPS = ["owner", "manager", "supervisor", "crew"]
MGMT = ["owner", "manager", "supervisor"]
PERM = {
    "pos": OPS, "menu_edit": ["owner", "manager"], "inventory_view": OPS,
    "inventory_move": MGMT, "opname": OPS, "waste": OPS, "waste_approve": MGMT,
    "purchase": OPS, "purchase_approve": ["owner", "manager"], "purchase_receive": MGMT,
    "finance": ["owner", "finance"], "payroll": ["owner", "finance"],
    "attendance_all": ["owner", "manager", "supervisor", "finance"],
    "marketing": ["owner", "marketing", "manager"], "users": ["owner"], "settings": ["owner", "finance"],
    "sales_view": ["owner", "manager", "supervisor", "finance", "marketing", "crew"],
    "sales_void": ["owner", "manager", "supervisor"],
    "customers": ["owner", "manager", "supervisor", "marketing", "crew"],
    "attendance_live": ["owner", "manager", "supervisor", "finance", "crew"],
    "overtime_approve": ["owner", "manager", "supervisor"],
    "audit": ["owner", "finance"],
    "petty_cash": OPS,
    "petty_cash_approve": MGMT + ["finance"],
    "adjust": ["owner", "manager", "supervisor", "finance"],
    "external_sales": OPS + ["finance"],
}
REPORTS = {
    "sales": ["owner", "manager", "supervisor", "finance", "marketing"],
    "nota": ["owner", "manager", "supervisor", "finance"],
    "inventory": MGMT, "opname": MGMT, "waste": MGMT, "purchasing": ["owner", "manager", "supervisor", "finance"],
    "finance": ["owner", "finance"], "payroll": ["owner", "finance"],
    "attendance": ["owner", "manager", "supervisor", "finance"], "marketing": ["owner", "marketing", "manager"],
}


def now():
    return datetime.now(TZ)


def today():
    return now().strftime("%Y-%m-%d")


def nid():
    return str(uuid.uuid4())


def hash_pw(p):
    return bcrypt.hashpw(p.encode(), bcrypt.gensalt()).decode()


def verify_pw(p, h):
    return bcrypt.checkpw(p.encode(), h.encode())


def make_token(u):
    t = datetime.now(timezone.utc)
    return jwt.encode({"sub": u["id"], "role": u["role"], "iat": t, "exp": t + timedelta(hours=12)}, JWT_SECRET, algorithm="HS256")


def public_user(u):
    out = {k: u.get(k) for k in ["id", "email", "name", "role", "position", "base_salary", "allowance", "is_active", "outlet_id",
                                  "hourly_rate", "overtime_rate", "shift_start"]}
    out["has_pin"] = bool(u.get("pin_hash"))
    return out


async def current_user(cred: HTTPAuthorizationCredentials = Depends(bearer)):
    if not cred:
        raise HTTPException(401, "Belum login")
    try:
        payload = jwt.decode(cred.credentials, JWT_SECRET, algorithms=["HS256"])
    except jwt.InvalidTokenError:
        raise HTTPException(401, "Token tidak valid / kedaluwarsa")
    u = await db.users.find_one({"id": payload["sub"]}, {"_id": 0})
    if not u or not u.get("is_active", True):
        raise HTTPException(401, "User tidak aktif")
    return u


def need(perm):
    async def dep(u=Depends(current_user)):
        if u["role"] not in PERM[perm]:
            raise HTTPException(403, "Akses ditolak untuk jabatan Anda")
        return u
    return dep


async def find(col, q=None, sort="created_at", limit=1000):
    q = {**(q or {}), "outlet_id": OUTLET}
    return await db[col].find(q, {"_id": 0}).sort(sort, -1).to_list(limit)


async def insert(col, doc):
    doc = {"id": nid(), "outlet_id": OUTLET, "created_at": now().isoformat(), "date": today(), **doc}
    await db[col].insert_one(dict(doc))
    return doc


# ---------------- Inventory core ----------------
async def get_stock(ing_id, loc):
    s = await db.stocks.find_one({"ingredient_id": ing_id, "location": loc, "outlet_id": OUTLET}, {"_id": 0})
    return s["qty"] if s else 0.0


async def move_stock(ing_id, loc, delta, mtype, ref, user):
    await db.stocks.update_one({"ingredient_id": ing_id, "location": loc, "outlet_id": OUTLET},
                               {"$inc": {"qty": delta}}, upsert=True)
    ing = await db.ingredients.find_one({"id": ing_id}, {"_id": 0})
    await insert("stock_movements", {"ingredient_id": ing_id, "ingredient_name": ing["name"] if ing else "",
                                     "unit": ing["unit"] if ing else "", "location": loc, "qty": delta,
                                     "type": mtype, "ref": ref, "user_name": user.get("name", "")})


async def ingredient_map():
    return {i["id"]: i for i in await find("ingredients", sort="name")}


async def menu_hpp(menu, imap):
    return round(sum(r["qty"] * imap.get(r["ingredient_id"], {}).get("cost_per_unit", 0) for r in menu.get("recipe", [])), 2)


def stock_location(menu_loc, st):
    """Simple mode -> semua stok di satu pool (gudang). Detailed -> per lokasi menu."""
    return "gudang" if (st or {}).get("stock_mode") == "simple" else (menu_loc or "bar")


def sale_total(x):
    """Nilai akhir nota (pakai final_total jika ada penyesuaian pembayaran)."""
    return x.get("final_total", x.get("total", 0))


async def resolve_pocket(category):
    c = await db.expense_categories.find_one({"name": category, "outlet_id": OUTLET}, {"_id": 0})
    if c and c.get("pocket_id"):
        return c["pocket_id"]
    p = await db.fund_pockets.find_one({"name": category, "outlet_id": OUTLET}, {"_id": 0})
    return p["id"] if p else None


async def add_finance(ftype, category, amount, note, ref, user, **extra):
    """Single cash/finance ledger. Income/expense lines; negative amounts = reversals (never delete)."""
    doc = {"type": ftype, "category": category, "amount": round(amount, 2), "note": note, "ref": ref,
           "user_name": user.get("name", ""), "user_id": user.get("id"), "source": "sistem", "status": "approved",
           "method": None, "pocket_id": None, "session_id": None, **extra}
    if ftype == "expense" and not doc["pocket_id"]:
        doc["pocket_id"] = await resolve_pocket(category)
    return await insert("finance", doc)


async def audit(entity, entity_id, action, user, before=None, after=None, reason=""):
    """Append-only audit trail (no delete endpoint exists)."""
    await insert("audit_logs", {"entity": entity, "entity_id": entity_id, "action": action, "before": before,
                                "after": after, "reason": reason, "user_name": user.get("name", ""),
                                "user_id": user.get("id"), "role": user.get("role")})


async def payment_fee(method, amount):
    pm = await db.payment_methods.find_one({"name": method, "outlet_id": OUTLET}, {"_id": 0})
    if not pm:
        return 0.0
    if pm.get("kind") == "qris":
        q = (await get_settings()).get("qris", {})
        return round(amount * float(q.get("mdr_percent") or 0) / 100 + float(q.get("fixed_fee") or 0), 2)
    return round(amount * float(pm.get("fee_percent") or 0) / 100 + float(pm.get("fee_fixed") or 0), 2)


# ---------------- Models ----------------
class LoginIn(BaseModel):
    email: str
    password: str


class UserIn(BaseModel):
    email: str
    password: Optional[str] = None
    name: str
    role: str
    position: str = ""
    base_salary: float = 0
    allowance: float = 0
    is_active: bool = True
    hourly_rate: float = 0
    overtime_rate: float = 0
    shift_start: str = ""  # HH:MM scheduled start, used for lateness
    pin: Optional[str] = None


class PayRateIn(BaseModel):
    base_salary: float = 0
    allowance: float = 0
    hourly_rate: float = 0
    overtime_rate: float = 0
    shift_start: str = ""


class IngredientIn(BaseModel):
    name: str
    unit: str
    cost_per_unit: float
    min_stock: float = 0
    category: str = "Bahan"


class RecipeItem(BaseModel):
    ingredient_id: str
    qty: float


class MenuIn(BaseModel):
    name: str
    category: str
    price: float
    location: str = "bar"
    station: str = "bar"  # bar | kitchen -> which order ticket prints the item
    recipe: List[RecipeItem] = []
    is_active: bool = True


class SaleItem(BaseModel):
    menu_id: Optional[str] = None  # None = item manual (bukan dari katalog)
    name: str = ""
    price: Optional[float] = None
    qty: int
    note: str = ""


class SaleIn(BaseModel):
    items: List[SaleItem]
    payment_method: str
    promo_id: Optional[str] = None
    customer: str = ""
    customer_id: Optional[str] = None
    apply_tax: bool = True
    cash_received: float = 0
    order_note: str = ""


class MoveIn(BaseModel):
    ingredient_id: str
    type: str  # in | out | transfer
    qty: float
    location: str
    to_location: Optional[str] = None
    note: str = ""
    total_cost: float = 0


class OpnameLine(BaseModel):
    ingredient_id: str
    physical: float


class OpnameIn(BaseModel):
    location: str
    lines: List[OpnameLine]
    note: str = ""


class WasteIn(BaseModel):
    ingredient_id: str
    qty: float
    reason: str
    location: str


class PRItem(BaseModel):
    ingredient_id: str
    qty: float
    est_price: float = 0


class PRIn(BaseModel):
    items: List[PRItem]
    note: str = ""


class ReceiveIn(BaseModel):
    total_paid: Optional[float] = None


class DecisionIn(BaseModel):
    approve: bool
    note: str = ""


class FinanceIn(BaseModel):
    type: str  # income | expense
    category: str
    amount: float
    note: str = ""


class AttendanceIn(BaseModel):
    type: str  # clock_in | clock_out | izin | sakit | libur | lembur
    note: str = ""
    hours: float = 0
    photo: str = ""  # base64 bukti foto kamera (wajib untuk clock in/out dari UI)


class SlipComponent(BaseModel):
    component_id: str
    qty: float = 1
    amount: Optional[float] = None  # only for method "manual"
    note: str = ""


class PayrollIn(BaseModel):
    user_id: str
    month: str  # YYYY-MM
    components: List[SlipComponent] = []


class PromoIn(BaseModel):
    name: str
    kind: str = "promo"  # promo | campaign
    discount_type: str = "percent"  # percent | nominal
    value: float = 0
    start_date: str
    end_date: str
    description: str = ""
    budget: float = 0
    is_active: bool = True


class ExpenseIn(BaseModel):
    category: str
    amount: float
    note: str = ""
    method: str = "Cash"
    pocket_id: Optional[str] = None
    receiver: str = ""
    ref: str = ""
    photo: str = ""  # base64 bukti
    source: Optional[str] = None


class AdjustIn(BaseModel):
    type: str  # refund | additional
    amount: float
    method: Optional[str] = None
    note: str = ""


class PrintLogIn(BaseModel):
    station: str = "customer"
    status: str = "success"  # success | failed
    note: str = ""


class ExtItem(BaseModel):
    menu_id: Optional[str] = None
    name: str = ""
    qty: int = 1
    price: float = 0
    note: str = ""


class ExtSaleIn(BaseModel):
    platform: str
    order_no: str = ""
    items: List[ExtItem]
    discount: float = 0
    discount_borne_by: Optional[str] = None  # outlet | platform
    commission_percent: Optional[float] = None
    other_fee_percent: Optional[float] = None
    fee_fixed: Optional[float] = None
    payment_status: str = "pending"  # pending | settled
    customer: str = ""
    note: str = ""


# ---------------- Auth & users ----------------
@api.post("/auth/login")
async def login(data: LoginIn):
    u = await db.users.find_one({"email": data.email.lower().strip()}, {"_id": 0})
    if not u or not verify_pw(data.password, u["password_hash"]) or not u.get("is_active", True):
        raise HTTPException(401, "Email atau password salah")
    return {"access_token": make_token(u), "user": public_user(u)}


class SessionIn(BaseModel):
    session_id: str


@api.post("/auth/session")
async def google_session(data: SessionIn):
    """Emergent-managed Google sign-in: exchange one-time session_id, then issue ROLLOW JWT.
    Only emails registered by the Owner may sign in (role comes from our DB)."""
    async with httpx.AsyncClient(timeout=15) as hc:
        r = await hc.get("https://demobackend.emergentagent.com/auth/v1/env/oauth/session-data",
                         headers={"X-Session-ID": data.session_id})
    if r.status_code != 200:
        raise HTTPException(401, "Sesi Google tidak valid")
    email = r.json().get("email", "").lower()
    u = await db.users.find_one({"email": email}, {"_id": 0})
    if not u or not u.get("is_active", True):
        raise HTTPException(403, f"Email {email} belum terdaftar. Minta Owner menambahkan Anda di menu Karyawan.")
    token = make_token(u)
    return {"session_token": token, "access_token": token, "user": public_user(u)}


@api.get("/auth/me")
async def me(u=Depends(current_user)):
    return public_user(u)


@api.get("/users")
async def list_users(u=Depends(current_user)):
    if u["role"] not in PERM["users"] + PERM["payroll"] + PERM["attendance_all"]:
        raise HTTPException(403, "Akses ditolak")
    return [public_user(x) for x in await find("users", sort="name")]


@api.post("/users")
async def create_user(data: UserIn, u=Depends(need("users"))):
    if data.role not in ALL:
        raise HTTPException(400, "Role tidak valid")
    if await db.users.find_one({"email": data.email.lower()}):
        raise HTTPException(409, "Email sudah terdaftar")
    if not data.password or len(data.password) < 6:
        raise HTTPException(400, "Password minimal 6 karakter")
    d = data.model_dump()
    d["email"] = d["email"].lower()
    d["password_hash"] = hash_pw(d.pop("password"))
    pin = d.pop("pin")
    if pin:
        d["pin_hash"] = hash_pw(pin)
    return public_user(await insert("users", d))


@api.put("/users/{uid}")
async def update_user(uid: str, data: UserIn, u=Depends(need("users"))):
    d = data.model_dump()
    pw = d.pop("password")
    pin = d.pop("pin")
    d["email"] = d["email"].lower()
    if pw:
        d["password_hash"] = hash_pw(pw)
    if pin:
        if not pin.isdigit() or not 4 <= len(pin) <= 6:
            raise HTTPException(400, "PIN harus 4-6 angka")
        d["pin_hash"] = hash_pw(pin)
    await db.users.update_one({"id": uid}, {"$set": d})
    return public_user(await db.users.find_one({"id": uid}, {"_id": 0}))


@api.put("/users/{uid}/pay")
async def update_pay(uid: str, data: PayRateIn, u=Depends(need("payroll"))):
    await db.users.update_one({"id": uid, "outlet_id": OUTLET}, {"$set": data.model_dump()})
    return public_user(await db.users.find_one({"id": uid}, {"_id": 0}))


async def verify_pin(user_id, pin):
    emp = await db.users.find_one({"id": user_id, "outlet_id": OUTLET}, {"_id": 0})
    if not emp or not emp.get("is_active", True):
        raise HTTPException(404, "Karyawan tidak ditemukan")
    if not emp.get("pin_hash"):
        raise HTTPException(400, f"{emp['name']} belum punya PIN. Minta Owner mengatur PIN di menu Karyawan.")
    if not verify_pw(pin or "", emp["pin_hash"]):
        raise HTTPException(400, "PIN salah")
    return emp


# ---------------- Ingredients & inventory ----------------
@api.get("/ingredients")
async def list_ingredients(u=Depends(current_user)):
    ings = await find("ingredients", sort="name")
    stocks = await find("stocks", sort="location")
    for i in ings:
        i["stock"] = {loc: 0.0 for loc in LOCATIONS}
        for s in stocks:
            if s["ingredient_id"] == i["id"]:
                i["stock"][s["location"]] = round(s["qty"], 2)
        i["total"] = round(sum(i["stock"].values()), 2)
        i["low"] = i["total"] < i.get("min_stock", 0)
    return sorted(ings, key=lambda x: x["name"])


@api.post("/ingredients")
async def create_ingredient(data: IngredientIn, u=Depends(need("menu_edit"))):
    return await insert("ingredients", data.model_dump())


@api.put("/ingredients/{iid}")
async def update_ingredient(iid: str, data: IngredientIn, u=Depends(need("menu_edit"))):
    await db.ingredients.update_one({"id": iid}, {"$set": data.model_dump()})
    return {"ok": True}


@api.get("/inventory/low-stock")
async def low_stock(u=Depends(current_user)):
    return [i for i in await list_ingredients(u) if i["low"]]


@api.post("/inventory/move")
async def inventory_move(data: MoveIn, u=Depends(need("inventory_move"))):
    if data.qty <= 0 or data.location not in LOCATIONS:
        raise HTTPException(400, "Data tidak valid")
    if data.type == "in":
        await move_stock(data.ingredient_id, data.location, data.qty, "stok_masuk", data.note, u)
        if data.total_cost > 0:
            await add_finance("expense", "Purchasing", data.total_cost, f"Stok masuk: {data.note}", "", u)
    elif data.type == "out":
        await move_stock(data.ingredient_id, data.location, -data.qty, "stok_keluar", data.note, u)
    elif data.type == "transfer":
        if data.to_location not in LOCATIONS or data.to_location == data.location:
            raise HTTPException(400, "Lokasi tujuan tidak valid")
        if await get_stock(data.ingredient_id, data.location) < data.qty:
            raise HTTPException(400, "Stok lokasi asal tidak cukup")
        await move_stock(data.ingredient_id, data.location, -data.qty, "transfer_keluar", f"ke {data.to_location}", u)
        await move_stock(data.ingredient_id, data.to_location, data.qty, "transfer_masuk", f"dari {data.location}", u)
    else:
        raise HTTPException(400, "Tipe tidak valid")
    return {"ok": True}


@api.get("/inventory/movements")
async def movements(u=Depends(need("inventory_view"))):
    return await find("stock_movements", limit=200)


# ---------------- Menu ----------------
@api.get("/menus")
async def list_menus(u=Depends(current_user)):
    imap = await ingredient_map()
    menus = await find("menus", sort="name")
    for m in menus:
        m["hpp"] = await menu_hpp(m, imap)
        m["margin"] = round(m["price"] - m["hpp"], 2)
        m["recipe_detail"] = [{**r, "name": imap.get(r["ingredient_id"], {}).get("name", "?"),
                               "unit": imap.get(r["ingredient_id"], {}).get("unit", "")} for r in m.get("recipe", [])]
    return sorted(menus, key=lambda x: (x["category"], x["name"]))


@api.post("/menus")
async def create_menu(data: MenuIn, u=Depends(need("menu_edit"))):
    return await insert("menus", data.model_dump())


@api.put("/menus/{mid}")
async def update_menu(mid: str, data: MenuIn, u=Depends(need("menu_edit"))):
    await db.menus.update_one({"id": mid}, {"$set": data.model_dump()})
    return {"ok": True}


# ---------------- Promo ----------------
async def active_promos():
    t = today()
    return [p for p in await find("promos", {"is_active": True}) if p["start_date"] <= t <= p["end_date"] and p["kind"] == "promo"]


@api.get("/promos")
async def list_promos(active: bool = False, u=Depends(current_user)):
    if active:
        return await active_promos()
    promos = await find("promos")
    for p in promos:
        sales = await find("sales", {"promo_id": p["id"], **VALID_SALE})
        p["usage"] = len(sales)
        p["revenue"] = round(sum(s["total"] for s in sales), 2)
        p["discount_given"] = round(sum(s["discount"] for s in sales), 2)
    return promos


@api.post("/promos")
async def create_promo(data: PromoIn, u=Depends(need("marketing"))):
    p = await insert("promos", data.model_dump())
    if data.budget > 0:
        await add_finance("expense", "Marketing", data.budget, f"Budget {data.kind}: {data.name}", p["id"], u)
    return p


@api.put("/promos/{pid}")
async def update_promo(pid: str, data: PromoIn, u=Depends(need("marketing"))):
    await db.promos.update_one({"id": pid}, {"$set": data.model_dump(exclude={"budget"})})
    return {"ok": True}


# ---------------- Settings (tax, loyalty, receipt) ----------------
DEFAULT_SETTINGS = {
    "tax_enabled": True, "tax_name": "PB1", "tax_percent": 10.0, "tax_default_on": True,
    "service_percent": 0.0,
    "late_tolerance_min": 10, "late_penalty_fixed": 0.0, "late_penalty_per_minute": 0.0,
    "loyalty_enabled": True, "points_per_nota": 1, "points_min_purchase": 0, "points_per_amount": 0,
    "petty_cash_limit": 50000.0,
    "stock_mode": "detailed",  # detailed = gudang/bar/kitchen terpisah | simple = satu stok total
    "close_report_detail": True,  # True = laporan tutup kasir rinci (per metode, void, pajak) | False = ringkas
    "petty_cash_needs_approval": False,
    "printer": {"paper_width": 58},
    "qris": {"provider": "", "merchant_name": "", "merchant_id": "", "bank_name": "", "account_no": "", "account_name": "",
             "settlement_info": "", "mdr_percent": 0.7, "fixed_fee": 0, "effective_date": ""},
    "receipt": {"cafe_name": "RLW", "address": "", "phone": "", "header_note": "",
                "footer_note": "Terima kasih atas kunjungan Anda!", "wifi": "",
                "show_cashier": True, "show_customer": True, "show_points": True},
}


class SettingsIn(BaseModel):
    tax_enabled: bool
    tax_name: str
    tax_percent: float
    tax_default_on: bool
    service_percent: float = 0
    loyalty_enabled: bool
    points_per_nota: int
    points_min_purchase: float = 0
    points_per_amount: float = 0
    late_tolerance_min: int = 10
    late_penalty_fixed: float = 0
    late_penalty_per_minute: float = 0
    petty_cash_limit: float = 50000
    stock_mode: str = "detailed"
    close_report_detail: bool = True
    petty_cash_needs_approval: bool = False
    printer: dict = {}
    qris: dict = {}
    receipt: dict


async def get_settings():
    d = await db.settings.find_one({"outlet_id": OUTLET}, {"_id": 0})
    out = {**DEFAULT_SETTINGS, **(d or {})}
    for k in ["receipt", "printer", "qris"]:
        out[k] = {**DEFAULT_SETTINGS[k], **((d or {}).get(k) or {})}
    return out


@api.get("/settings")
async def read_settings(u=Depends(current_user)):
    st = await get_settings()
    if u["role"] not in PERM["settings"]:  # hide sensitive bank / merchant data
        st["qris"] = {k: st["qris"].get(k) for k in ["provider", "merchant_name", "mdr_percent", "fixed_fee"]}
    return st


@api.put("/settings")
async def write_settings(data: SettingsIn, u=Depends(need("settings"))):
    await audit("settings", OUTLET, "update", u, before=await get_settings(), after=data.model_dump())
    await db.settings.update_one({"outlet_id": OUTLET}, {"$set": {**data.model_dump(), "outlet_id": OUTLET,
                                 "updated_by": u["name"], "updated_at": now().isoformat()}}, upsert=True)
    return await get_settings()


# ---------------- Customers / loyalty ----------------
class CustomerIn(BaseModel):
    name: str
    phone: str = ""


@api.get("/customers")
async def list_customers(q: str = "", u=Depends(need("customers"))):
    query = {"$or": [{"name": {"$regex": q, "$options": "i"}}, {"phone": {"$regex": q}}]} if q else {}
    return await find("customers", query, sort="points", limit=200)


@api.post("/customers")
async def create_customer(data: CustomerIn, u=Depends(need("pos"))):
    if data.phone and await db.customers.find_one({"phone": data.phone, "outlet_id": OUTLET}):
        raise HTTPException(409, "No. HP sudah terdaftar")
    return await insert("customers", {**data.model_dump(), "points": 0, "visits": 0, "total_spent": 0})


@api.get("/customers/{cid}/history")
async def customer_history(cid: str, u=Depends(need("customers"))):
    return await find("sales", {"customer_id": cid}, limit=100)


# ---------------- Register session & shifts (shared outlet tablet) ----------------
VALID_SALE = {"status": {"$nin": ["void", "cancelled"]}}


class OpenIn(BaseModel):
    opening_cash: float = 0


class HandoverIn(BaseModel):
    user_id: str
    pin: str
    note: str = ""


async def open_session():
    return await db.register_sessions.find_one({"outlet_id": OUTLET, "status": "open"}, {"_id": 0})


async def session_view(sess):
    if not sess:
        return None
    sales = await find("sales", {"session_id": sess["id"], **VALID_SALE}, limit=5000)
    for sh in sess["shifts"]:
        mine = [x for x in sales if x.get("shift_no") == sh["no"]]
        sh["count"] = len(mine)
        sh["total"] = round(sum(x["total"] for x in mine), 2)
        by = {}
        for x in mine:
            by[x["payment_method"]] = by.get(x["payment_method"], 0) + x["total"]
        sh["by_method"] = by
    sess["current"] = sess["shifts"][-1]
    return sess


@api.get("/register/current")
async def register_current(u=Depends(need("pos"))):
    return {"session": await session_view(await open_session())}


@api.post("/register/open")
async def register_open(data: OpenIn, u=Depends(need("pos"))):
    if await open_session():
        raise HTTPException(400, "Kasir sudah dibuka")
    t = now()
    sess = await insert("register_sessions", {"status": "open", "opening_cash": data.opening_cash, "opened_by": u["name"],
                                              "shifts": [{"no": 1, "user_id": u["id"], "user_name": u["name"],
                                                          "start": t.strftime("%H:%M"), "start_ts": t.isoformat(), "end": None}]})
    return {"session": await session_view(sess)}


@api.post("/register/handover")
async def register_handover(data: HandoverIn, u=Depends(need("pos"))):
    sess = await open_session()
    if not sess:
        raise HTTPException(400, "Kasir belum dibuka")
    emp = await verify_pin(data.user_id, data.pin)
    cur = sess["shifts"][-1]
    if cur["user_id"] == emp["id"]:
        raise HTTPException(400, "Karyawan ini sedang memegang shift")
    t = now()
    cur.update({"end": t.strftime("%H:%M"), "end_ts": t.isoformat(), "note": data.note})
    sess["shifts"].append({"no": cur["no"] + 1, "user_id": emp["id"], "user_name": emp["name"],
                           "start": t.strftime("%H:%M"), "start_ts": t.isoformat(), "end": None})
    await db.register_sessions.update_one({"id": sess["id"]}, {"$set": {"shifts": sess["shifts"]}})
    return {"session": await session_view(sess)}


# ---------------- POS / Sales ----------------


@api.post("/sales")
async def create_sale(data: SaleIn, u=Depends(need("pos"))):
    if not data.items:
        raise HTTPException(400, "Keranjang kosong")
    sess = await open_session()
    if not sess:
        raise HTTPException(400, "Kasir belum dibuka. Tekan 'Buka Kasir' dulu.")
    shift = sess["shifts"][-1]
    st = await get_settings()
    imap = await ingredient_map()
    lines, subtotal, hpp_total = [], 0.0, 0.0
    usage = {}
    for it in data.items:
        if it.menu_id:
            m = await db.menus.find_one({"id": it.menu_id}, {"_id": 0})
            if not m:
                raise HTTPException(404, "Menu tidak ditemukan")
            hpp = await menu_hpp(m, imap)
            lines.append({"menu_id": m["id"], "name": m["name"], "qty": it.qty, "price": m["price"], "note": it.note,
                          "station": m.get("station") or m.get("location", "bar"),
                          "hpp": hpp, "subtotal": m["price"] * it.qty, "manual": False})
            subtotal += m["price"] * it.qty
            hpp_total += hpp * it.qty
            for r in m.get("recipe", []):
                key = (r["ingredient_id"], stock_location(m.get("location", "bar"), st))
                usage[key] = usage.get(key, 0) + r["qty"] * it.qty
        else:
            # item manual (bukan dari katalog): tidak memotong stok & tanpa HPP
            if not str(it.name).strip() or it.price is None or it.price < 0:
                raise HTTPException(400, "Item manual butuh nama & harga")
            price = float(it.price)
            lines.append({"menu_id": None, "name": it.name.strip(), "qty": it.qty, "price": price, "note": it.note,
                          "station": "bar", "hpp": 0, "subtotal": price * it.qty, "manual": True})
            subtotal += price * it.qty
    discount, promo_name = 0.0, None
    if data.promo_id:
        p = await db.promos.find_one({"id": data.promo_id}, {"_id": 0})
        if p:
            discount = subtotal * p["value"] / 100 if p["discount_type"] == "percent" else min(p["value"], subtotal)
            promo_name = p["name"]
    base = subtotal - discount
    service = round(base * st["service_percent"] / 100, 2) if st["service_percent"] else 0.0
    tax = round((base + service) * st["tax_percent"] / 100, 2) if (st["tax_enabled"] and data.apply_tax) else 0.0
    total = round(base + service + tax, 2)
    cust, points = None, 0
    if data.customer_id:
        cust = await db.customers.find_one({"id": data.customer_id}, {"_id": 0})
        if cust and st["loyalty_enabled"] and total >= st["points_min_purchase"]:
            points = int(st["points_per_nota"]) + (int(total // st["points_per_amount"]) if st["points_per_amount"] else 0)
    count = await db.sales.count_documents({"date": today(), "outlet_id": OUTLET})
    sale = await insert("sales", {
        "number": f"RLW-{now().strftime('%y%m%d')}-{count + 1:04d}", "items": lines, "subtotal": subtotal,
        "discount": round(discount, 2), "promo_id": data.promo_id if promo_name else None, "promo_name": promo_name,
        "service": service, "tax": tax, "tax_name": st["tax_name"] if tax else None, "tax_percent": st["tax_percent"] if tax else 0,
        "total": total, "hpp": round(hpp_total, 2), "profit": round(base - hpp_total, 2),
        "payment_method": data.payment_method, "cash_received": data.cash_received,
        "customer": cust["name"] if cust else data.customer, "customer_id": cust["id"] if cust else None,
        "customer_phone": cust["phone"] if cust else "", "points_earned": points,
        "points_balance": (cust["points"] + points) if cust else None, "order_note": data.order_note,
        "time": now().strftime("%H:%M"), "status": "paid",
        "session_id": sess["id"], "shift_no": shift["no"], "input_by": u["name"],
        "cashier_id": shift["user_id"], "cashier_name": shift["user_name"], "closed": False})
    # recipe -> auto deduct inventory
    for (iid, loc), q in usage.items():
        await move_stock(iid, loc, -q, "penjualan", sale["number"], u)
    # sale -> finance income (tax collected recorded separately as liability-like income line)
    tag = {"method": data.payment_method, "session_id": sess["id"], "source": "kasir"}
    await add_finance("income", "Penjualan", round(base + service, 2), f"Penjualan {sale['number']} ({data.payment_method})", sale["id"], u, **tag)
    if tax:
        await add_finance("income", f"Pajak {st['tax_name']}", tax, f"Pajak {sale['number']}", sale["id"], u, **tag)
    fee = await payment_fee(data.payment_method, total)
    await db.sales.update_one({"id": sale["id"]}, {"$set": {"payment_fee": fee, "expected_net": round(total - fee, 2)}})
    sale.update({"payment_fee": fee, "expected_net": round(total - fee, 2)})
    if fee:
        await add_finance("expense", "Biaya Pembayaran", fee, f"MDR/biaya {data.payment_method} {sale['number']}", sale["id"], u,
                          method=data.payment_method, source="sistem")
    if cust:
        await db.customers.update_one({"id": cust["id"]}, {"$inc": {"points": points, "visits": 1, "total_spent": total}})
    return sale


class VoidIn(BaseModel):
    kind: str = "void"  # void (cancelled legacy masih diterima utk data lama)
    reason: str


@api.post("/sales/{sid}/void")
async def void_sale(sid: str, data: VoidIn, u=Depends(need("sales_void"))):
    s = await db.sales.find_one({"id": sid, "outlet_id": OUTLET}, {"_id": 0})
    if not s or s.get("status") in ["void", "cancelled"]:
        raise HTTPException(400, "Nota tidak ditemukan / sudah dibatalkan")
    if data.kind not in ["void", "cancelled"] or not data.reason.strip():
        raise HTTPException(400, "Isi alasan pembatalan")
    # reverse stock
    movs = await db.stock_movements.find({"ref": s["number"], "type": "penjualan", "outlet_id": OUTLET}, {"_id": 0}).to_list(500)
    for mv in movs:
        await move_stock(mv["ingredient_id"], mv["location"], -mv["qty"], f"{data.kind}_penjualan", s["number"], u)
    # void = transaksi batal total -> nilai kembali 0, BUKAN baris minus.
    # Baris pemasukan nota dinolkan (nilai asli disimpan untuk audit), tidak ada refund bernilai negatif.
    fees = await db.finance.find({"ref": sid, "type": "expense", "category": "Biaya Pembayaran", "outlet_id": OUTLET}, {"_id": 0}).to_list(5)
    for f in fees:
        if not f.get("voided"):
            await db.finance.update_one({"id": f["id"]}, {"$set": {"voided": True, "original_amount": f["amount"], "amount": 0,
                                                                    "void_note": f"Batal biaya {s['number']}: {data.reason}"}})
    incomes = await db.finance.find({"ref": sid, "type": "income", "outlet_id": OUTLET}, {"_id": 0}).to_list(50)
    for fr in incomes:
        if not fr.get("voided"):
            await db.finance.update_one({"id": fr["id"]}, {"$set": {"voided": True, "original_amount": fr["amount"], "amount": 0,
                                                                     "void_note": f"Void {s['number']}: {data.reason}"}})
    refund = await insert("refunds", {"sale_id": sid, "number": s["number"], "amount": s["total"], "method": s["payment_method"],
                                      "reason": data.reason, "status": "paid", "paid_by": u["name"], "paid_at": now().isoformat(),
                                      "created_by": u["name"]})
    # reverse loyalty
    if s.get("customer_id"):
        await db.customers.update_one({"id": s["customer_id"]}, {"$inc": {"points": -s.get("points_earned", 0), "visits": -1, "total_spent": -s["total"]}})
    await db.sales.update_one({"id": sid}, {"$set": {"status": data.kind, "void_reason": data.reason,
                                                    "void_by": u["name"], "void_at": now().isoformat()}})
    await audit("sale", sid, data.kind, u, before={"status": s.get("status", "paid"), "total": s["total"]},
                after={"status": data.kind, "refund_id": refund["id"]}, reason=data.reason)
    return {"ok": True, "status": data.kind, "refund": refund}


@api.get("/sales")
async def list_sales(date: Optional[str] = None, period: str = "day", status: str = "all", u=Depends(need("sales_view"))):
    d = date or today()
    q = {"date": {"$regex": f"^{d if period == 'day' else d[:7]}"}}
    if status != "all":
        q["status"] = status if status != "paid" else {"$nin": ["void", "cancelled"]}
    sales = await find("sales", q, limit=2000)
    valid = [x for x in sales if x.get("status", "paid") not in ["void", "cancelled"]]
    by, by_channel = {}, {}
    for x in valid:
        by[x["payment_method"]] = by.get(x["payment_method"], 0) + sale_total(x)
        ch = x.get("channel", "kasir")
        by_channel[ch] = by_channel.get(ch, 0) + sale_total(x)
    recap = {"nota": len(valid), "void": len([x for x in sales if x.get("status") == "void"]),
             "cancelled": len([x for x in sales if x.get("status") == "cancelled"]),
             "subtotal": sum(x["subtotal"] for x in valid), "discount": sum(x["discount"] for x in valid),
             "tax": sum(x.get("tax", 0) for x in valid), "service": sum(x.get("service", 0) for x in valid),
             "total": sum(sale_total(x) for x in valid), "hpp": sum(x["hpp"] for x in valid),
             "profit": sum(x["profit"] for x in valid), "by_method": by, "by_channel": by_channel}
    return {"recap": recap, "sales": sales}


@api.get("/pos/summary")
async def pos_summary(u=Depends(need("pos"))):
    sess = await session_view(await open_session())
    sales = await find("sales", {"session_id": sess["id"], **VALID_SALE}, limit=5000) if sess else []
    voided = await find("sales", {"session_id": sess["id"], "status": {"$in": ["void", "cancelled"]}}, limit=5000) if sess else []
    by, by_detail = {}, {}
    for x in sales:
        pm = x["payment_method"]
        by[pm] = by.get(pm, 0) + x["total"]
        d = by_detail.setdefault(pm, {"count": 0, "amount": 0.0})
        d["count"] += 1
        d["amount"] = round(d["amount"] + x["total"], 2)
    gross = round(sum(x["subtotal"] for x in sales), 2)
    discount = round(sum(x["discount"] for x in sales), 2)
    tax = round(sum(x.get("tax", 0) for x in sales), 2)
    service = round(sum(x.get("service", 0) for x in sales), 2)
    net = round(sum(x["total"] for x in sales), 2)
    return {"count": len(sales), "total": net,
            "discount": discount, "by_method": by, "by_method_detail": by_detail,
            "gross": gross, "tax": tax, "service": service, "net": net,
            "void_count": len(voided), "void_amount": round(sum(x["total"] for x in voided), 2),
            "opening_cash": sess["opening_cash"] if sess else 0, "shifts": sess["shifts"] if sess else [],
            "cash_ledger": await cash_ledger(sess),
            "closings": await find("register_closes", {"date": today()})}


async def cash_ledger(sess):
    """Rekonsiliasi laci kasir — CASH ONLY. Hanya mutasi tunai (method 'Cash')
    yang terikat sesi berjalan yang memengaruhi laci. Transaksi non-tunai diabaikan."""
    if not sess:
        return {"cash_in": 0, "cash_out": 0, "expected_cash": 0,
                "kas_awal": 0, "total_pemasukan_cash": 0, "total_pengeluaran_cash": 0, "kas_sistem": 0}
    rows = await db.finance.find({"session_id": sess["id"], "method": "Cash", "outlet_id": OUTLET}, {"_id": 0}).to_list(5000)
    cin = sum(r["amount"] for r in rows if r["type"] == "income")
    cout = sum(r["amount"] for r in rows if r["type"] == "expense")
    kas_awal = sess.get("opening_cash", 0) or 0
    kas_sistem = round(kas_awal + cin - cout, 2)
    return {"cash_in": round(cin, 2), "cash_out": round(cout, 2), "expected_cash": kas_sistem,
            "kas_awal": round(kas_awal, 2), "total_pemasukan_cash": round(cin, 2),
            "total_pengeluaran_cash": round(cout, 2), "kas_sistem": kas_sistem}


@api.get("/cash/outlet")
async def cash_outlet(u=Depends(need("petty_cash"))):
    """Kas fisik outlet (realtime): saldo awal + uang masuk tunai − pengeluaran kas outlet.
    Hanya menghitung mutasi tunai (method Cash) yang terikat sesi kasir berjalan.
    Pengeluaran Finance (sumber dana perusahaan) TIDAK mengurangi saldo ini."""
    sess = await open_session()
    if not sess:
        return {"open": False, "opening_cash": 0, "cash_in": 0, "cash_out": 0, "balance": 0, "history": []}
    rows = await db.finance.find({"session_id": sess["id"], "method": "Cash", "outlet_id": OUTLET},
                                 {"_id": 0}).sort("created_at", -1).to_list(5000)
    cin = sum(r["amount"] for r in rows if r["type"] == "income")
    cout = sum(r["amount"] for r in rows if r["type"] == "expense")
    hist = [{"id": r["id"], "type": r["type"], "category": r.get("category", ""), "amount": round(r["amount"], 2),
             "note": r.get("note", ""), "user_name": r.get("user_name", ""), "date": r.get("date"),
             "time": (r.get("created_at") or "")[11:16], "created_at": r.get("created_at")} for r in rows if not r.get("voided")]
    return {"open": True, "opened_by": sess.get("opened_by"), "opened_at": sess.get("created_at"),
            "opening_cash": round(sess["opening_cash"], 2), "cash_in": round(cin, 2), "cash_out": round(cout, 2),
            "balance": round(sess["opening_cash"] + cin - cout, 2), "history": hist}


class CloseIn(BaseModel):
    kas_aktual: float = 0
    selisih_note: str = ""


@api.post("/pos/close")
async def pos_close(data: CloseIn, u=Depends(need("pos"))):
    sess = await open_session()
    if not sess:
        raise HTTPException(400, "Kasir belum dibuka")
    summ = await pos_summary(u)
    summ.pop("closings")
    ledger = summ["cash_ledger"]
    kas_sistem = ledger["kas_sistem"]
    kas_aktual = round(data.kas_aktual or 0, 2)
    selisih = round(kas_aktual - kas_sistem, 2)
    note = (data.selisih_note or "").strip()
    if selisih != 0 and not note:
        raise HTTPException(400, "Kas tidak balance — wajib isi alasan selisih")
    t = now()
    shifts = summ["shifts"]
    shifts[-1].update({"end": t.strftime("%H:%M"), "end_ts": t.isoformat()})
    rec = await insert("register_closes", {**summ, "session_id": sess["id"], "user_name": u["name"],
                                           "expected_cash": kas_sistem,
                                           "kas_awal": ledger["kas_awal"],
                                           "total_pemasukan_cash": ledger["total_pemasukan_cash"],
                                           "total_pengeluaran_cash": ledger["total_pengeluaran_cash"],
                                           "kas_sistem": kas_sistem, "kas_aktual": kas_aktual,
                                           "selisih": selisih, "selisih_note": note})
    await db.register_sessions.update_one({"id": sess["id"]}, {"$set": {"status": "closed", "shifts": shifts, "closed_by": u["name"], "closed_at": t.isoformat()}})
    await db.sales.update_many({"session_id": sess["id"]}, {"$set": {"closed": True}})
    return rec


RECAP_ROLES = ["owner", "manager", "supervisor", "finance", "marketing"]


@api.get("/sales/recap")
async def sales_recap(period: str = "day", date: Optional[str] = None, u=Depends(current_user)):
    """Rekap penjualan simpel & sinkron: omzet, selisih kas, top 5 terlaris, jam ramai."""
    if u["role"] not in RECAP_ROLES:
        raise HTTPException(403, "Akses ditolak untuk jabatan Anda")
    p = rng(period, date)
    q = {"date": {"$regex": f"^{p}"}}
    allsales = await find("sales", q, limit=10000)
    valid = [x for x in allsales if x.get("status", "paid") not in ["void", "cancelled"]]
    omzet = sum(sale_total(x) for x in valid)
    prod = {}
    for s in valid:
        for i in s["items"]:
            d = prod.setdefault(i["name"], {"name": i["name"], "qty": 0, "omzet": 0.0})
            d["qty"] += i["qty"]
            d["omzet"] += i.get("subtotal", 0)
    top = sorted(prod.values(), key=lambda x: (-x["qty"], -x["omzet"]))[:5]
    for d in top:
        d["omzet"] = round(d["omzet"], 2)
    hours = {}
    for s in valid:
        hh = (s.get("time") or (s.get("created_at") or "")[11:16])[:2]
        if hh.isdigit():
            h = int(hh)
            d = hours.setdefault(h, {"hour": h, "count": 0, "omzet": 0.0})
            d["count"] += 1
            d["omzet"] += sale_total(s)
    busy = sorted(hours.values(), key=lambda x: x["hour"])
    for d in busy:
        d["omzet"] = round(d["omzet"], 2)
    peak = max(busy, key=lambda x: x["count"]) if busy else None
    by = {}
    for s in valid:
        by[s["payment_method"]] = round(by.get(s["payment_method"], 0) + sale_total(s), 2)
    closes = await find("register_closes", q, limit=500)
    kas = {"closes": len(closes),
           "kas_sistem": round(sum(c.get("kas_sistem", 0) for c in closes), 2),
           "kas_aktual": round(sum(c.get("kas_aktual", 0) for c in closes), 2),
           "selisih": round(sum(c.get("selisih", 0) for c in closes), 2)}
    close_rows = [{"user_name": c.get("user_name"), "kas_sistem": round(c.get("kas_sistem", 0), 2),
                   "kas_aktual": round(c.get("kas_aktual", 0), 2), "selisih": round(c.get("selisih", 0), 2),
                   "selisih_note": c.get("selisih_note", ""), "time": (c.get("created_at") or "")[11:16],
                   "date": c.get("date")} for c in closes]
    return {"period": period, "range": p,
            "trx": len(valid), "void": len([x for x in allsales if x.get("status") == "void"]),
            "cancelled": len([x for x in allsales if x.get("status") == "cancelled"]),
            "omzet": round(omzet, 2), "discount": round(sum(x.get("discount", 0) for x in valid), 2),
            "tax": round(sum(x.get("tax", 0) for x in valid), 2), "hpp": round(sum(x.get("hpp", 0) for x in valid), 2),
            "profit": round(sum(x.get("profit", 0) for x in valid), 2),
            "by_method": [{"method": k, "total": v} for k, v in sorted(by.items(), key=lambda x: -x[1])],
            "top_products": top, "busy_hours": busy, "peak_hour": peak, "kas": kas, "closes_detail": close_rows}


# ---------------- Stock opname ----------------
@api.post("/opname")
async def create_opname(data: OpnameIn, u=Depends(need("opname"))):
    imap = await ingredient_map()
    lines = []
    for ln in data.lines:
        system = await get_stock(ln.ingredient_id, data.location)
        diff = round(ln.physical - system, 2)
        ing = imap.get(ln.ingredient_id, {})
        lines.append({"ingredient_id": ln.ingredient_id, "name": ing.get("name"), "unit": ing.get("unit"),
                      "system": round(system, 2), "physical": ln.physical, "diff": diff,
                      "diff_value": round(diff * ing.get("cost_per_unit", 0), 2)})
        if diff != 0:
            await move_stock(ln.ingredient_id, data.location, diff, "selisih_opname", "Stock opname", u)
    return await insert("opnames", {"location": data.location, "lines": lines, "note": data.note,
                                    "total_diff_value": round(sum(x["diff_value"] for x in lines), 2),
                                    "user_name": u["name"]})


@api.get("/opname")
async def list_opname(u=Depends(need("opname"))):
    return await find("opnames", limit=100)


# ---------------- Waste ----------------
@api.post("/waste")
async def create_waste(data: WasteIn, u=Depends(need("waste"))):
    ing = await db.ingredients.find_one({"id": data.ingredient_id}, {"_id": 0})
    if not ing or data.qty <= 0:
        raise HTTPException(400, "Data tidak valid")
    return await insert("wastes", {**data.model_dump(), "ingredient_name": ing["name"], "unit": ing["unit"],
                                   "value": round(data.qty * ing["cost_per_unit"], 2), "status": "pending",
                                   "user_name": u["name"]})


@api.get("/waste")
async def list_waste(u=Depends(need("waste"))):
    return await find("wastes", limit=200)


@api.post("/waste/{wid}/decide")
async def decide_waste(wid: str, data: DecisionIn, u=Depends(need("waste_approve"))):
    w = await db.wastes.find_one({"id": wid}, {"_id": 0})
    if not w or w["status"] != "pending":
        raise HTTPException(400, "Waste tidak ditemukan / sudah diproses")
    status = "approved" if data.approve else "rejected"
    await db.wastes.update_one({"id": wid}, {"$set": {"status": status, "approved_by": u["name"]}})
    if data.approve:
        await move_stock(w["ingredient_id"], w["location"], -w["qty"], "waste", w["reason"], u)
    return {"ok": True, "status": status}


# ---------------- Purchasing ----------------
@api.post("/purchasing")
async def create_pr(data: PRIn, u=Depends(need("purchase"))):
    imap = await ingredient_map()
    items = [{**i.model_dump(), "name": imap.get(i.ingredient_id, {}).get("name"),
              "unit": imap.get(i.ingredient_id, {}).get("unit"),
              "est_price": i.est_price or round(i.qty * imap.get(i.ingredient_id, {}).get("cost_per_unit", 0), 2)}
             for i in data.items]
    if not items:
        raise HTTPException(400, "Item kosong")
    count = await db.purchase_requests.count_documents({"outlet_id": OUTLET})
    return await insert("purchase_requests", {"number": f"PR-{count + 1:04d}", "items": items, "note": data.note,
                                              "est_total": round(sum(i["est_price"] for i in items), 2),
                                              "status": "pending", "requested_by": u["name"]})


@api.get("/purchasing")
async def list_pr(u=Depends(need("purchase"))):
    return await find("purchase_requests", limit=200)


@api.post("/purchasing/{pid}/decide")
async def decide_pr(pid: str, data: DecisionIn, u=Depends(need("purchase_approve"))):
    pr = await db.purchase_requests.find_one({"id": pid}, {"_id": 0})
    if not pr or pr["status"] != "pending":
        raise HTTPException(400, "PR tidak ditemukan / sudah diproses")
    status = "approved" if data.approve else "rejected"
    await db.purchase_requests.update_one({"id": pid}, {"$set": {"status": status, "approved_by": u["name"], "decision_note": data.note}})
    return {"ok": True, "status": status}


@api.post("/purchasing/{pid}/receive")
async def receive_pr(pid: str, data: ReceiveIn, u=Depends(need("purchase_receive"))):
    pr = await db.purchase_requests.find_one({"id": pid}, {"_id": 0})
    if not pr or pr["status"] != "approved":
        raise HTTPException(400, "PR belum disetujui")
    for i in pr["items"]:
        await move_stock(i["ingredient_id"], "gudang", i["qty"], "purchasing", pr["number"], u)
    paid = data.total_paid if data.total_paid is not None else pr["est_total"]
    await add_finance("expense", "Purchasing", paid, f"Pembelian {pr['number']}", pr["id"], u)
    await db.purchase_requests.update_one({"id": pid}, {"$set": {"status": "received", "total_paid": paid, "received_by": u["name"]}})
    return {"ok": True}


# ---------------- Finance ----------------
@api.post("/finance")
async def create_finance(data: FinanceIn, u=Depends(need("finance"))):
    if data.type not in ["income", "expense"] or data.amount <= 0:
        raise HTTPException(400, "Data tidak valid")
    return await add_finance(data.type, data.category, data.amount, data.note, "", u)


@api.get("/finance")
async def list_finance(month: Optional[str] = None, u=Depends(need("finance"))):
    m = month or today()[:7]
    return await db.finance.find({"outlet_id": OUTLET, "date": {"$regex": f"^{m}"}}, {"_id": 0}).sort("created_at", -1).to_list(2000)


# ---------------- Attendance ----------------
ATT_TYPES = ["clock_in", "clock_out", "izin", "sakit", "libur", "lembur"]


class KioskIn(BaseModel):
    user_id: str
    pin: str
    type: str
    note: str = ""
    hours: float = 0
    photo: str = ""


def hm_to_min(hm):
    try:
        h, m = hm.split(":")
        return int(h) * 60 + int(m)
    except Exception:
        return None


async def record_attendance(emp, type_, note, hours, by_name, photo=""):
    if type_ not in ATT_TYPES:
        raise HTTPException(400, "Tipe tidak valid")
    mine = await find("attendance", {"user_id": emp["id"], "date": today()})
    types = [a["type"] for a in mine]
    if type_ == "clock_in" and "clock_in" in types:
        raise HTTPException(400, f"{emp['name']} sudah clock in hari ini")
    if type_ == "clock_out" and ("clock_in" not in types or "clock_out" in types):
        raise HTTPException(400, f"{emp['name']} belum clock in / sudah clock out")
    if type_ == "lembur" and hours <= 0:
        raise HTTPException(400, "Isi jumlah jam lembur")
    t = now()
    extra = {}
    if type_ == "clock_in":
        st = await get_settings()
        sched = hm_to_min(emp.get("shift_start") or "")
        late = max(0, t.hour * 60 + t.minute - sched) if sched is not None else 0
        extra = {"late_minutes": late, "is_late": late > st["late_tolerance_min"], "shift_start": emp.get("shift_start")}
    if type_ == "lembur":
        extra = {"status": "pending"}
    return await insert("attendance", {"type": type_, "note": note, "hours": hours, "user_id": emp["id"], "user_name": emp["name"],
                                       "role": emp["role"], "time": t.strftime("%H:%M"), "ts": t.isoformat(),
                                       "recorded_by": by_name, "photo": photo or "", **extra})


@api.post("/attendance")
async def create_attendance(data: AttendanceIn, u=Depends(current_user)):
    return await record_attendance(u, data.type, data.note, data.hours, u["name"], data.photo)


@api.post("/attendance/kiosk")
async def kiosk_attendance(data: KioskIn, u=Depends(need("pos"))):
    emp = await verify_pin(data.user_id, data.pin)
    return await record_attendance(emp, data.type, data.note, data.hours, f"Tab outlet ({u['name']})", data.photo)


@api.get("/attendance/staff")
async def kiosk_staff(u=Depends(need("pos"))):
    users = await find("users", {"is_active": True}, sort="name")
    return [{"id": x["id"], "name": x["name"], "role": x["role"], "position": x.get("position", ""), "has_pin": bool(x.get("pin_hash"))}
            for x in sorted(users, key=lambda x: x["name"]) if x["role"] != "owner"]


@api.get("/attendance/live")
async def attendance_live(u=Depends(need("attendance_live"))):
    users = [x for x in await find("users", {"is_active": True}) if x["role"] != "owner"]
    att = await find("attendance", {"date": today()})
    t = now()
    rows = []
    for emp in sorted(users, key=lambda x: x["name"]):
        mine = sorted([a for a in att if a["user_id"] == emp["id"]], key=lambda a: a["created_at"])
        cin = next((a for a in mine if a["type"] == "clock_in"), None)
        cout = next((a for a in mine if a["type"] == "clock_out"), None)
        other = next((a for a in mine if a["type"] in ["izin", "sakit", "libur"]), None)
        hours = 0.0
        if cin:
            end = datetime.fromisoformat(cout["created_at"]) if cout else t
            hours = round((end - datetime.fromisoformat(cin["created_at"])).total_seconds() / 3600, 2)
        status = "bertugas" if cin and not cout else "selesai" if cout else (other["type"] if other else "belum hadir")
        rows.append({"user_id": emp["id"], "name": emp["name"], "role": emp["role"], "status": status,
                     "clock_in": cin["time"] if cin else None, "clock_out": cout["time"] if cout else None,
                     "late_minutes": cin.get("late_minutes", 0) if cin else 0, "is_late": cin.get("is_late", False) if cin else False,
                     "clock_in_photo": cin.get("photo") if cin else None, "clock_out_photo": cout.get("photo") if cout else None,
                     "hours": hours, "shift_start": emp.get("shift_start") or "",
                     "overtime_pending": sum(a["hours"] for a in mine if a["type"] == "lembur" and a.get("status") == "pending")})
    summary = {"bertugas": len([r for r in rows if r["status"] == "bertugas"]), "selesai": len([r for r in rows if r["status"] == "selesai"]),
               "terlambat": len([r for r in rows if r["is_late"]]), "belum_hadir": len([r for r in rows if r["status"] == "belum hadir"]),
               "izin_sakit_libur": len([r for r in rows if r["status"] in ["izin", "sakit", "libur"]]),
               "total": len(rows), "time": t.strftime("%H:%M:%S")}
    return {"summary": summary, "rows": rows}


@api.post("/attendance/{aid}/decide")
async def decide_overtime(aid: str, data: DecisionIn, u=Depends(need("overtime_approve"))):
    a = await db.attendance.find_one({"id": aid, "outlet_id": OUTLET}, {"_id": 0})
    if not a or a["type"] != "lembur" or a.get("status") != "pending":
        raise HTTPException(400, "Lembur tidak ditemukan / sudah diproses")
    status = "approved" if data.approve else "rejected"
    await db.attendance.update_one({"id": aid}, {"$set": {"status": status, "approved_by": u["name"]}})
    return {"ok": True, "status": status}


@api.get("/attendance")
async def list_attendance(all: bool = False, month: Optional[str] = None, pending: bool = False, u=Depends(current_user)):
    q = {"date": {"$regex": f"^{month or today()[:7]}"}}
    if pending:
        q = {"type": "lembur", "status": "pending"}
    if not ((all or pending) and u["role"] in PERM["attendance_all"]):
        q["user_id"] = u["id"]
    return await find("attendance", q, limit=1000)


# ---------------- Payroll (auto from attendance) ----------------
async def compute_pay(emp, month, st):
    att = sorted(await find("attendance", {"user_id": emp["id"], "date": {"$regex": f"^{month}"}}), key=lambda a: a["created_at"])
    t = now()
    days = {}
    for a in att:
        days.setdefault(a["date"], []).append(a)
    hours, present, late_count, late_minutes, running = 0.0, 0, 0, 0, False
    for d, recs in days.items():
        cin = next((a for a in recs if a["type"] == "clock_in"), None)
        cout = next((a for a in recs if a["type"] == "clock_out"), None)
        if not cin:
            continue
        present += 1
        if cin.get("is_late"):
            late_count += 1
            late_minutes += cin.get("late_minutes", 0)
        if cout:
            hours += (datetime.fromisoformat(cout["created_at"]) - datetime.fromisoformat(cin["created_at"])).total_seconds() / 3600
        elif d == today():
            running = True
            hours += (t - datetime.fromisoformat(cin["created_at"])).total_seconds() / 3600
    hours = round(hours, 2)
    ot = [a for a in att if a["type"] == "lembur"]
    ot_ok = sum(a["hours"] for a in ot if a.get("status", "approved") == "approved")
    ot_pending = sum(a["hours"] for a in ot if a.get("status") == "pending")
    hourly = emp.get("hourly_rate", 0) or 0
    ot_rate = emp.get("overtime_rate", 0) or hourly
    late_penalty = late_count * st["late_penalty_fixed"] + late_minutes * st["late_penalty_per_minute"]
    hourly_pay = round(hours * hourly, 2)
    overtime_pay = round(ot_ok * ot_rate, 2)
    gross = emp.get("base_salary", 0) + emp.get("allowance", 0) + hourly_pay + overtime_pay
    return {"user_id": emp["id"], "name": emp["name"], "role": emp["role"], "month": month,
            "base_salary": emp.get("base_salary", 0), "allowance": emp.get("allowance", 0),
            "hourly_rate": hourly, "overtime_rate": ot_rate, "shift_start": emp.get("shift_start") or "",
            "hours": hours, "running": running, "hourly_pay": hourly_pay, "overtime_hours": ot_ok,
            "overtime_pending": ot_pending, "overtime": overtime_pay, "late_count": late_count, "late_minutes": late_minutes,
            "late_penalty": round(late_penalty, 2),
            "attendance": {"clock_in": present, **{k: len([a for a in att if a["type"] == k]) for k in ["izin", "sakit", "libur"]},
                           "lembur_hours": ot_ok},
            "estimate": round(gross - late_penalty, 2)}


@api.get("/payroll/preview")
async def payroll_preview(month: Optional[str] = None, u=Depends(need("payroll"))):
    m = month or today()[:7]
    st = await get_settings()
    users = [x for x in await find("users", {"is_active": True}) if x["role"] != "owner"]
    rows = [await compute_pay(e, m, st) for e in sorted(users, key=lambda x: x["name"])]
    return {"month": m, "rows": rows, "total": round(sum(r["estimate"] for r in rows), 2),
            "hours": round(sum(r["hours"] for r in rows), 2), "late_penalty": round(sum(r["late_penalty"] for r in rows), 2)}


@api.post("/payroll")
async def create_payroll(data: PayrollIn, u=Depends(need("payroll"))):
    emp = await db.users.find_one({"id": data.user_id}, {"_id": 0})
    if not emp:
        raise HTTPException(404, "Karyawan tidak ditemukan")
    if await db.payrolls.find_one({"user_id": data.user_id, "month": data.month, "outlet_id": OUTLET}):
        raise HTTPException(409, "Payroll bulan ini sudah dibuat")
    c = await compute_pay(emp, data.month, await get_settings())
    c.pop("estimate")
    bonuses, deductions = [], []
    for sc in data.components:
        comp = await db.payroll_components.find_one({"id": sc.component_id, "outlet_id": OUTLET}, {"_id": 0})
        if not comp or not comp.get("is_active", True):
            raise HTTPException(400, "Komponen gaji tidak valid / nonaktif")
        m, a = comp["method"], float(comp.get("amount") or 0)
        val = {"fixed": a * sc.qty, "percent_base": c["base_salary"] * a / 100, "per_late": c["late_count"] * a,
               "per_late_minute": c["late_minutes"] * a, "per_day_present": c["attendance"]["clock_in"] * a,
               "per_hour_worked": c["hours"] * a, "manual": sc.amount or 0}.get(m, 0)
        line = {"component_id": comp["id"], "name": comp["name"], "description": comp.get("description", ""), "method": m,
                "rate": a, "qty": sc.qty, "amount": round(val, 2), "note": sc.note}
        (bonuses if comp["type"] == "bonus" else deductions).append(line)
    tb = round(sum(x["amount"] for x in bonuses), 2)
    td = round(sum(x["amount"] for x in deductions), 2)
    gross = c["base_salary"] + c["allowance"] + c["hourly_pay"] + c["overtime"]
    # snapshot: later component changes never alter a saved slip
    return await insert("payrolls", {**c, "bonuses": bonuses, "deductions": deductions, "total_bonus": tb, "total_deduction": td,
                                     "bonus": tb, "deduction": td, "late_penalty": 0, "gross": round(gross, 2),
                                     "total": round(gross + tb - td, 2), "status": "draft", "created_by": u["name"]})


@api.get("/payroll")
async def list_payroll(month: Optional[str] = None, u=Depends(need("payroll"))):
    return await find("payrolls", {"month": month or today()[:7]})


@api.post("/payroll/{pid}/pay")
async def pay_payroll(pid: str, u=Depends(need("payroll"))):
    p = await db.payrolls.find_one({"id": pid}, {"_id": 0})
    if not p or p["status"] == "paid":
        raise HTTPException(400, "Payroll tidak valid / sudah dibayar")
    await add_finance("expense", "Gaji", p["total"], f"Gaji {p['name']} {p['month']}", pid, u, source="payroll", method="Transfer")
    await audit("payroll", pid, "pay", u, before={"status": p["status"]}, after={"status": "paid", "total": p["total"]})
    await db.payrolls.update_one({"id": pid}, {"$set": {"status": "paid", "paid_at": now().isoformat()}})
    return {"ok": True}


# ---------------- Dashboard ----------------
@api.get("/dashboard")
async def dashboard(u=Depends(current_user)):
    r, t, m = u["role"], today(), today()[:7]
    out = {"role": r, "date": t}
    if r in OPS + ["finance", "marketing"]:
        sales = await find("sales", {"date": t, **VALID_SALE})
        out["sales_today"] = round(sum(s["total"] for s in sales), 2)
        out["trx_today"] = len(sales)
        out["profit_today"] = round(sum(s["profit"] for s in sales), 2)
    if r in OPS:
        low = await low_stock(u)
        out["low_stock"] = [{"name": i["name"], "total": i["total"], "min": i["min_stock"], "unit": i["unit"]} for i in low]
    if r in MGMT:
        out["pending_waste"] = await db.wastes.count_documents({"status": "pending", "outlet_id": OUTLET})
        out["pending_pr"] = await db.purchase_requests.count_documents({"status": "pending", "outlet_id": OUTLET})
        out["approved_pr"] = await db.purchase_requests.count_documents({"status": "approved", "outlet_id": OUTLET})
    if r in ["owner", "finance"]:
        fin = await list_finance(m, u)
        inc = sum(f["amount"] for f in fin if f["type"] == "income")
        exp = sum(f["amount"] for f in fin if f["type"] == "expense")
        out["finance_month"] = {"income": round(inc, 2), "expense": round(exp, 2), "net": round(inc - exp, 2)}
    if r in PERM["marketing"]:
        out["active_promos"] = len(await active_promos())
    if r in PERM["overtime_approve"]:
        out["pending_overtime"] = await db.attendance.count_documents({"type": "lembur", "status": "pending", "outlet_id": OUTLET})
    mine = await find("attendance", {"user_id": u["id"], "date": t})
    out["my_attendance"] = [a["type"] for a in mine]
    return out


# ---------------- Reports ----------------
def rng(period, date):
    d = date or today()
    return d if period == "day" else d[:7]


# ---------------- Excel export (openpyxl) ----------------
_HDR_FILL = PatternFill("solid", fgColor="1F6F54")
_HDR_FONT = Font(bold=True, color="FFFFFF")
_TITLE_FONT = Font(bold=True, size=14, color="1F6F54")
_TOTAL_FONT = Font(bold=True)
_THIN = Side(style="thin", color="D0D5DD")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _sheet_header(ws, headers, row_idx):
    for ci, h in enumerate(headers, start=1):
        c = ws.cell(row=row_idx, column=ci, value=h)
        c.fill = _HDR_FILL
        c.font = _HDR_FONT
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = _BORDER


def _autosize(ws):
    from openpyxl.cell.cell import MergedCell
    widths = {}
    for row in ws.iter_rows():
        for c in row:
            if isinstance(c, MergedCell) or c.value is None:
                continue
            widths[c.column_letter] = max(widths.get(c.column_letter, 10), len(str(c.value)))
    for letter, w in widths.items():
        ws.column_dimensions[letter].width = min(max(w + 2, 11), 42)


def _title_block(ws, cafe, label, period_label, ncols):
    ws.cell(row=1, column=1, value=f"{cafe} — {label}").font = _TITLE_FONT
    ws.cell(row=2, column=1, value=f"Periode: {period_label}")
    ws.cell(row=3, column=1, value=f"Dibuat: {now().strftime('%d-%m-%Y %H:%M')}")
    if ncols > 1:
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncols)
        ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=ncols)


@api.get("/reports/{rtype}/export")
async def export_report(rtype: str, period: str = "day", date: Optional[str] = None, u=Depends(current_user)):
    if rtype not in ("sales", "finance", "nota"):
        raise HTTPException(404, "Unduhan Excel tidak tersedia untuk laporan ini")
    if u["role"] not in ("owner", "finance"):
        raise HTTPException(403, "Hanya Owner & Finance yang dapat mengunduh laporan")
    p = rng(period, date)
    q = {"date": {"$regex": f"^{p}"}}
    st = await get_settings()
    cafe = (st.get("receipt") or {}).get("cafe_name") or "ROLLOW"
    label = {"sales": "Rekap Penjualan", "finance": "Laporan Keuangan", "nota": "Rekap per Nota"}[rtype]
    wb = Workbook()
    wb.remove(wb.active)

    if rtype == "sales":
        allsales = await find("sales", q, limit=20000)
        valid = [s for s in allsales if s.get("status", "paid") not in ("void", "cancelled")]
        menus = await find("menus", limit=2000)
        mcat = {m["id"]: (m.get("category") or "", m.get("sub_category") or "") for m in menus}
        agg, tx = {}, {}
        for s in valid:
            for i in s["items"]:
                key = i["name"]
                a = agg.setdefault(key, {"name": key, "cat": mcat.get(i.get("menu_id"), ("", ""))[0],
                                         "sub": mcat.get(i.get("menu_id"), ("", ""))[1], "qty": 0, "omzet": 0.0, "cogs": 0.0})
                a["qty"] += i["qty"]
                a["omzet"] += i.get("subtotal", 0)
                a["cogs"] += i.get("hpp", 0)
                tx.setdefault(key, set()).add(s["number"])
        tot_qty = sum(a["qty"] for a in agg.values()) or 1
        tot_omzet = sum(a["omzet"] for a in agg.values()) or 1
        # Sheet 1: Rekap Menu
        ws = wb.create_sheet("Rekap Menu")
        _title_block(ws, cafe, "Rekap Penjualan per Menu", p, 9)
        heads = ["Menu", "Kategori", "Sub Kategori", "Qty Terjual", "% Qty", "Penjualan", "% Penjualan", "COGS", "Margin"]
        hr = 5
        _sheet_header(ws, heads, hr)
        r = hr + 1
        for a in sorted(agg.values(), key=lambda x: -x["omzet"]):
            margin = round(a["omzet"] - a["cogs"], 2)
            vals = [a["name"], a["cat"], a["sub"], a["qty"], round(a["qty"] / tot_qty * 100, 2),
                    round(a["omzet"], 2), round(a["omzet"] / tot_omzet * 100, 2), round(a["cogs"], 2), margin]
            for ci, v in enumerate(vals, start=1):
                ws.cell(row=r, column=ci, value=v).border = _BORDER
            r += 1
        for ci, v in enumerate(["TOTAL", "", "", sum(a["qty"] for a in agg.values()), 100.0,
                                round(sum(a["omzet"] for a in agg.values()), 2), 100.0,
                                round(sum(a["cogs"] for a in agg.values()), 2),
                                round(sum(a["omzet"] - a["cogs"] for a in agg.values()), 2)], start=1):
            c = ws.cell(row=r, column=ci, value=v)
            c.font = _TOTAL_FONT
            c.border = _BORDER
        _autosize(ws)
        # Sheet 2: Detail Transaksi
        ws2 = wb.create_sheet("Detail Transaksi")
        _title_block(ws2, cafe, "Detail Transaksi Penjualan", p, 15)
        heads2 = ["No Transaksi", "Tanggal", "Jam", "Kasir", "Pelanggan", "Mode", "Metode Pembayaran", "Status",
                  "Subtotal", "Diskon", "Pajak", "Service", "Total", "HPP", "Laba"]
        _sheet_header(ws2, heads2, 5)
        r = 6
        for s in sorted(allsales, key=lambda x: (x.get("date", ""), x.get("time", ""))):
            vals = [s.get("number", ""), s.get("date", ""), s.get("time", ""), s.get("cashier_name", ""),
                    s.get("customer") or "-", s.get("channel", "kasir"), s.get("payment_method", ""),
                    s.get("status", "paid").upper(), round(s.get("subtotal", 0), 2), round(s.get("discount", 0), 2),
                    round(s.get("tax", 0), 2), round(s.get("service", 0), 2), round(s.get("total", 0), 2),
                    round(s.get("hpp", 0), 2), round(s.get("profit", 0), 2)]
            for ci, v in enumerate(vals, start=1):
                ws2.cell(row=r, column=ci, value=v).border = _BORDER
            r += 1
        _autosize(ws2)

    elif rtype == "finance":
        fin = await find("finance", q, limit=20000)
        inc = sum(f["amount"] for f in fin if f["type"] == "income")
        exp = sum(f["amount"] for f in fin if f["type"] == "expense")
        ws = wb.create_sheet("Ringkasan")
        _title_block(ws, cafe, "Laporan Keuangan — Ringkasan", p, 2)
        rows = [("Total Pemasukan", round(inc, 2)), ("Total Pengeluaran", round(exp, 2)),
                ("Cash Flow Bersih", round(inc - exp, 2))]
        _sheet_header(ws, ["Keterangan", "Nominal"], 5)
        r = 6
        for k, v in rows:
            ws.cell(row=r, column=1, value=k).border = _BORDER
            ws.cell(row=r, column=2, value=v).border = _BORDER
            r += 1
        ws.cell(row=5 + 1, column=1).font = _TOTAL_FONT
        # per-category breakdown
        cats = {}
        for f in fin:
            key = (f["type"], f.get("category", "Lainnya"))
            cats[key] = round(cats.get(key, 0) + f["amount"], 2)
        ws.cell(row=r + 1, column=1, value="Rincian per Kategori").font = _TOTAL_FONT
        _sheet_header(ws, ["Tipe", "Kategori", "Nominal"], r + 2)
        rr = r + 3
        for (t, c), v in sorted(cats.items(), key=lambda x: (x[0][0], -x[1])):
            ws.cell(row=rr, column=1, value="Pemasukan" if t == "income" else "Pengeluaran").border = _BORDER
            ws.cell(row=rr, column=2, value=c).border = _BORDER
            ws.cell(row=rr, column=3, value=v).border = _BORDER
            rr += 1
        _autosize(ws)
        # detail sheet
        ws2 = wb.create_sheet("Detail Transaksi")
        _title_block(ws2, cafe, "Detail Mutasi Keuangan", p, 8)
        _sheet_header(ws2, ["Tanggal", "Jam", "Tipe", "Kategori", "Metode", "Catatan", "User", "Nominal"], 5)
        r = 6
        for f in sorted(fin, key=lambda x: x.get("created_at", "")):
            vals = [f.get("date", ""), (f.get("created_at") or "")[11:16],
                    "Pemasukan" if f["type"] == "income" else "Pengeluaran", f.get("category", ""),
                    f.get("method") or "-", f.get("note", ""), f.get("user_name", ""), round(f["amount"], 2)]
            for ci, v in enumerate(vals, start=1):
                ws2.cell(row=r, column=ci, value=v).border = _BORDER
            r += 1
        _autosize(ws2)

    elif rtype == "nota":
        allsales = await find("sales", q, limit=20000)
        ws = wb.create_sheet("Rekap per Nota")
        _title_block(ws, cafe, "Rekap per Nota", p, 12)
        _sheet_header(ws, ["No Nota", "Tanggal", "Jam", "Status", "Metode", "Pelanggan", "Kasir",
                           "Jml Item", "Subtotal", "Diskon", "Pajak", "Total"], 5)
        r = 6
        for s in sorted(allsales, key=lambda x: (x.get("date", ""), x.get("time", ""))):
            vals = [s.get("number", ""), s.get("date", ""), s.get("time", ""), s.get("status", "paid").upper(),
                    s.get("payment_method", ""), s.get("customer") or "-", s.get("cashier_name", ""),
                    sum(i["qty"] for i in s.get("items", [])), round(s.get("subtotal", 0), 2),
                    round(s.get("discount", 0), 2), round(s.get("tax", 0), 2), round(s.get("total", 0), 2)]
            for ci, v in enumerate(vals, start=1):
                ws.cell(row=r, column=ci, value=v).border = _BORDER
            r += 1
        _autosize(ws)

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    fname = f"{label.replace(' ', '')}_{p}.xlsx"
    return StreamingResponse(buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": f'attachment; filename="{fname}"'})


@api.get("/reports/{rtype}")
async def report(rtype: str, period: str = "day", date: Optional[str] = None, u=Depends(current_user)):
    if rtype not in REPORTS:
        raise HTTPException(404, "Laporan tidak ada")
    if u["role"] not in REPORTS[rtype]:
        raise HTTPException(403, "Akses ditolak")
    p = rng(period, date)
    q = {"date": {"$regex": f"^{p}"}}
    S, rows = {}, []
    if rtype == "sales":
        sales = await find("sales", {**q, **VALID_SALE})
        items = {}
        for s in sales:
            for i in s["items"]:
                it = items.setdefault(i["name"], {"qty": 0, "omzet": 0})
                it["qty"] += i["qty"]
                it["omzet"] += i["subtotal"]
        omzet = sum(s["total"] for s in sales)
        hpp = sum(s["hpp"] for s in sales)
        S = {"Transaksi": len(sales), "Omzet": omzet, "Diskon": sum(s["discount"] for s in sales),
             "Pajak": sum(s.get("tax", 0) for s in sales), "HPP": hpp, "Laba Kotor": omzet - hpp - sum(s.get("tax", 0) for s in sales),
             "Void/Batal": await db.sales.count_documents({**q, "outlet_id": OUTLET, "status": {"$in": ["void", "cancelled"]}})}
        for pm in set(s["payment_method"] for s in sales):
            S[f"Bayar {pm}"] = sum(s["total"] for s in sales if s["payment_method"] == pm)
        rows = [{"title": k, "subtitle": f"{v['qty']} terjual", "value": v["omzet"]} for k, v in sorted(items.items(), key=lambda x: -x[1]["omzet"])]
    elif rtype == "nota":
        allsales = await find("sales", q, limit=5000)
        valid = [x for x in allsales if x.get("status", "paid") not in ["void", "cancelled"]]
        S = {"Transaksi": len(valid), "Void/Batal": len(allsales) - len(valid), "Omzet": sum(x["total"] for x in valid),
             "Pajak": sum(x.get("tax", 0) for x in valid), "Diskon": sum(x["discount"] for x in valid)}
        rows = [{"title": f"{x['number']} · {x.get('time', '')}", "flag": x.get("status") in ["void", "cancelled"],
                 "subtitle": f"{x.get('status', 'paid').upper()} · {x['payment_method']} · {x.get('customer') or '-'} · {sum(i['qty'] for i in x['items'])} item · {x['cashier_name']}",
                 "value": x["total"]} for x in allsales]
    elif rtype == "inventory":
        ings = await list_ingredients(u)
        S = {"Total Bahan": len(ings), "Stok Rendah": len([i for i in ings if i["low"]]),
             "Nilai Stok": round(sum(i["total"] * i["cost_per_unit"] for i in ings), 2)}
        rows = [{"title": i["name"], "subtitle": f"G {i['stock']['gudang']} · B {i['stock']['bar']} · K {i['stock']['kitchen']} {i['unit']} (min {i['min_stock']})",
                 "value": round(i["total"] * i["cost_per_unit"], 2), "flag": i["low"]} for i in ings]
    elif rtype == "opname":
        ops = await find("opnames", q)
        agg = {}
        for o in ops:
            for ln in o["lines"]:
                a = agg.setdefault(ln["name"], {"diff": 0, "value": 0, "unit": ln["unit"]})
                a["diff"] += ln["diff"]
                a["value"] += ln["diff_value"]
        S = {"Jumlah Opname": len(ops), "Total Nilai Selisih": round(sum(o["total_diff_value"] for o in ops), 2)}
        rows = [{"title": k, "subtitle": f"Selisih {round(v['diff'], 2)} {v['unit']}", "value": round(v["value"], 2), "flag": v["diff"] < 0}
                for k, v in agg.items() if v["diff"] != 0]
    elif rtype == "waste":
        ws = await find("wastes", q)
        ap = [w for w in ws if w["status"] == "approved"]
        S = {"Total Catatan": len(ws), "Disetujui": len(ap), "Nilai Waste": round(sum(w["value"] for w in ap), 2)}
        rows = [{"title": w["ingredient_name"], "subtitle": f"{w['qty']} {w['unit']} · {w['reason']} · {w['location']} · {w['status']}", "value": w["value"]} for w in ws]
    elif rtype == "purchasing":
        prs = await find("purchase_requests", q)
        S = {"Total PR": len(prs), "Pending": len([x for x in prs if x["status"] == "pending"]),
             "Diterima": len([x for x in prs if x["status"] == "received"]),
             "Total Dibayar": round(sum(x.get("total_paid", 0) for x in prs), 2)}
        rows = [{"title": x["number"], "subtitle": f"{x['status']} · {x['requested_by']}", "value": x.get("total_paid", x["est_total"])} for x in prs]
    elif rtype == "finance":
        fin = await find("finance", q, limit=5000)
        inc = sum(f["amount"] for f in fin if f["type"] == "income")
        exp = sum(f["amount"] for f in fin if f["type"] == "expense")
        sales = await find("sales", {**q, **VALID_SALE})
        hpp = sum(s["hpp"] for s in sales)
        S = {"Pemasukan": inc, "Pengeluaran": exp, "Cash Flow Bersih": inc - exp, "HPP Terjual": hpp,
             "Laba Kotor (Omzet-HPP)": sum(s["total"] for s in sales) - hpp}
        cats = {}
        for f in fin:
            k = (f["type"], f["category"])
            cats[k] = cats.get(k, 0) + f["amount"]
        rows = [{"title": c, "subtitle": "Pemasukan" if t == "income" else "Pengeluaran", "value": v if t == "income" else -v} for (t, c), v in cats.items()]
    elif rtype == "payroll":
        ps = await find("payrolls", {"month": p[:7]})
        S = {"Karyawan": len(ps), "Total Gaji": sum(x["total"] for x in ps), "Sudah Dibayar": sum(x["total"] for x in ps if x["status"] == "paid")}
        rows = [{"title": x["name"], "subtitle": f"{x['role']} · {x['status']}", "value": x["total"]} for x in ps]
    elif rtype == "attendance":
        att = await find("attendance", q)
        S = {t: len([a for a in att if a["type"] == t]) for t in ["clock_in", "clock_out", "izin", "sakit", "libur", "lembur"]}
        rows = [{"title": a["user_name"], "subtitle": f"{a['date']} {a['time']} · {a['type']}", "value": None} for a in att]
    elif rtype == "marketing":
        promos = await list_promos(False, u)
        S = {"Promo/Campaign": len(promos), "Total Penggunaan": sum(x["usage"] for x in promos),
             "Omzet dari Promo": sum(x["revenue"] for x in promos), "Total Diskon": sum(x["discount_given"] for x in promos)}
        rows = [{"title": x["name"], "subtitle": f"{x['kind']} · {x['usage']}x dipakai · diskon {x['discount_given']}", "value": x["revenue"]} for x in promos]
    return {"type": rtype, "period": period, "range": p, "summary": S, "rows": rows}


# ---------------- Pengeluaran (petty cash kasir + finance) ----------------
@api.post("/expenses")
async def create_expense(data: ExpenseIn, u=Depends(need("petty_cash"))):
    if data.amount <= 0:
        raise HTTPException(400, "Nominal tidak valid")
    st = await get_settings()
    needs = bool(st.get("petty_cash_needs_approval")) and u["role"] not in PERM["petty_cash_approve"]
    sess = await open_session()
    source = data.source or ("kasir" if u["role"] in OPS else "finance")
    if needs:
        doc = await insert("expenses_pending", {"category": data.category, "amount": round(data.amount, 2), "note": data.note,
                                                "method": data.method, "pocket_id": data.pocket_id, "receiver": data.receiver,
                                                "ref": data.ref, "photo": data.photo, "source": source, "status": "pending",
                                                "user_name": u["name"], "user_id": u["id"],
                                                "session_id": sess["id"] if sess else None})
        await audit("expense", doc["id"], "request", u, after={"category": data.category, "amount": doc["amount"]}, reason=data.note)
        return {"pending": True, **doc}
    fin = await add_finance("expense", data.category, round(data.amount, 2), data.note, data.ref, u, method=data.method,
                            source=source, pocket_id=data.pocket_id, receiver=data.receiver, photo=data.photo,
                            session_id=sess["id"] if (sess and source == "kasir" and data.method == "Cash") else None)
    await audit("expense", fin["id"], "create", u, after={"category": data.category, "amount": fin["amount"], "method": data.method}, reason=data.note)
    return {"pending": False, **fin}


@api.get("/expenses")
async def list_expenses(month: Optional[str] = None, date: Optional[str] = None, pending: bool = False, u=Depends(need("petty_cash"))):
    if pending:
        return await find("expenses_pending", {"status": "pending"}, limit=500)
    rows = await db.finance.find({"type": "expense", "outlet_id": OUTLET, "date": {"$regex": f"^{date or month or today()[:7]}"}},
                                 {"_id": 0}).sort("created_at", -1).to_list(3000)
    if u["role"] not in PERM["finance"]:  # kasir/crew hanya lihat pengeluaran kasir
        rows = [r for r in rows if r.get("source") == "kasir"]
    return rows


@api.post("/expenses/{eid}/decide")
async def decide_expense(eid: str, data: DecisionIn, u=Depends(need("petty_cash_approve"))):
    p = await db.expenses_pending.find_one({"id": eid, "outlet_id": OUTLET}, {"_id": 0})
    if not p or p["status"] != "pending":
        raise HTTPException(400, "Pengeluaran tidak ditemukan / sudah diproses")
    status = "approved" if data.approve else "rejected"
    upd = {"status": status, "approved_by": u["name"], "approved_at": now().isoformat(), "decide_note": data.note}
    if data.approve:
        sess = await open_session()
        fin = await add_finance("expense", p["category"], p["amount"], p["note"], p.get("ref", ""), u, method=p.get("method"),
                                source=p.get("source", "kasir"), pocket_id=p.get("pocket_id"), receiver=p.get("receiver", ""),
                                photo=p.get("photo", ""), requested_by=p.get("user_name"),
                                session_id=sess["id"] if (sess and p.get("source", "kasir") == "kasir" and p.get("method") == "Cash") else None)
        upd["finance_id"] = fin["id"]
    await db.expenses_pending.update_one({"id": eid}, {"$set": upd})
    await audit("expense", eid, "decide", u, before={"status": "pending"}, after={"status": status}, reason=data.note)
    return {"ok": True, "status": status}


# ---------------- Penyesuaian pembayaran nota ----------------
@api.post("/sales/{sid}/adjust")
async def adjust_sale(sid: str, data: AdjustIn, u=Depends(need("adjust"))):
    s = await db.sales.find_one({"id": sid, "outlet_id": OUTLET}, {"_id": 0})
    if not s or s.get("status") in ["void", "cancelled"]:
        raise HTTPException(400, "Nota tidak ditemukan / sudah dibatalkan")
    if data.type not in ["refund", "additional"] or data.amount <= 0:
        raise HTTPException(400, "Data penyesuaian tidak valid")
    sess = await open_session()
    before = sale_total(s)
    method = data.method or s["payment_method"]
    if data.type == "refund":
        after = round(before - data.amount, 2)
        await add_finance("income", "Refund Penyesuaian", -round(data.amount, 2), f"Refund penyesuaian {s['number']}: {data.note}",
                          sid, u, method=method, source="kasir", session_id=sess["id"] if sess else None)
    else:
        after = round(before + data.amount, 2)
        await add_finance("income", "Penjualan", round(data.amount, 2), f"Tambahan bayar {s['number']}: {data.note}",
                          sid, u, method=method, source="kasir", session_id=sess["id"] if sess else None)
    adj = {"type": data.type, "amount": round(data.amount, 2), "method": method, "note": data.note,
           "by": u["name"], "at": now().isoformat(), "before_total": before, "after_total": after}
    await db.sales.update_one({"id": sid}, {"$push": {"adjustments": adj}, "$set": {"final_total": after}})
    await audit("sale", sid, "adjust", u, before={"total": before}, after={"total": after, "type": data.type, "amount": data.amount}, reason=data.note)
    return {"ok": True, "adjustment": adj, "final_total": after}


# ---------------- Log cetak / cetak ulang nota (tidak mengubah transaksi) ----------------
@api.post("/sales/{sid}/print-log")
async def print_log(sid: str, data: PrintLogIn, u=Depends(need("sales_view"))):
    s = await db.sales.find_one({"id": sid, "outlet_id": OUTLET}, {"_id": 0})
    if not s:
        raise HTTPException(404, "Nota tidak ditemukan")
    log = {"station": data.station, "status": data.status, "note": data.note, "by": u["name"], "at": now().isoformat()}
    await db.sales.update_one({"id": sid}, {"$push": {"print_logs": log}, "$inc": {"print_count": 1}})
    return {"ok": True, "log": log}


# ---------------- Penjualan platform eksternal (ShopeeFood/GoFood/GrabFood) ----------------
@api.post("/external-sales")
async def create_external_sale(data: ExtSaleIn, u=Depends(need("external_sales"))):
    if not data.items:
        raise HTTPException(400, "Minimal 1 item")
    plat = await db.platforms.find_one({"name": data.platform, "outlet_id": OUTLET}, {"_id": 0})
    if not plat:
        raise HTTPException(400, "Platform tidak dikenal")
    if data.order_no and await db.sales.find_one({"channel": "platform", "platform": data.platform, "order_no": data.order_no, "outlet_id": OUTLET}):
        raise HTTPException(409, "Nomor pesanan sudah pernah dicatat")
    st = await get_settings()
    imap = await ingredient_map()
    lines, subtotal, hpp_total, usage = [], 0.0, 0.0, {}
    for it in data.items:
        price, qty, name, hpp, station = float(it.price), int(it.qty), it.name, 0.0, "bar"
        if it.menu_id:
            m = await db.menus.find_one({"id": it.menu_id}, {"_id": 0})
            if m:
                name = name or m["name"]
                station = m.get("station") or m.get("location", "bar")
                hpp = await menu_hpp(m, imap)
                for r in m.get("recipe", []):
                    key = (r["ingredient_id"], stock_location(m.get("location", "bar"), st))
                    usage[key] = usage.get(key, 0) + r["qty"] * qty
        if not name:
            raise HTTPException(400, "Nama item wajib diisi")
        lines.append({"menu_id": it.menu_id, "name": name, "qty": qty, "price": price, "note": it.note,
                      "station": station, "hpp": hpp, "subtotal": price * qty})
        subtotal += price * qty
        hpp_total += hpp * qty
    gross = round(subtotal, 2)
    discount = round(float(data.discount or 0), 2)
    borne = data.discount_borne_by or plat.get("discount_borne_by", "outlet")
    comm_pct = data.commission_percent if data.commission_percent is not None else plat.get("commission_percent", 0)
    other_pct = data.other_fee_percent if data.other_fee_percent is not None else plat.get("other_fee_percent", 0)
    fee_fixed = data.fee_fixed if data.fee_fixed is not None else plat.get("fee_fixed", 0)
    commission = round(gross * float(comm_pct or 0) / 100, 2)
    other_fee = round(gross * float(other_pct or 0) / 100 + float(fee_fixed or 0), 2)
    outlet_disc = discount if borne == "outlet" else 0.0
    net = round(gross - commission - other_fee - outlet_disc, 2)
    count = await db.sales.count_documents({"date": today(), "outlet_id": OUTLET})
    sale = await insert("sales", {
        "number": f"EXT-{now().strftime('%y%m%d')}-{count + 1:04d}", "items": lines, "subtotal": gross,
        "discount": discount, "discount_borne_by": borne, "promo_id": None, "promo_name": None,
        "service": 0, "tax": 0, "tax_name": None, "tax_percent": 0, "total": gross, "final_total": gross,
        "hpp": round(hpp_total, 2), "profit": round(gross - hpp_total - commission - other_fee - outlet_disc, 2),
        "payment_method": data.platform, "channel": "platform", "platform": data.platform, "order_no": data.order_no,
        "commission": commission, "other_fee": other_fee, "platform_fee": round(commission + other_fee, 2), "expected_net": net,
        "cash_received": 0, "customer": data.customer, "customer_id": None, "customer_phone": "", "points_earned": 0,
        "order_note": data.note, "time": now().strftime("%H:%M"), "status": "paid", "payment_status": data.payment_status,
        "session_id": None, "shift_no": None, "input_by": u["name"], "cashier_id": u["id"], "cashier_name": u["name"], "closed": False})
    for (iid, loc), q in usage.items():
        await move_stock(iid, loc, -q, "penjualan_platform", sale["number"], u)
    tag = {"method": data.platform, "source": "platform", "channel": "platform"}
    if gross:
        await add_finance("income", "Penjualan", gross, f"Penjualan {data.platform} {sale['number']}", sale["id"], u, **tag)
    if commission:
        await add_finance("expense", "Komisi Platform", commission, f"Komisi {data.platform} {sale['number']}", sale["id"], u, method=data.platform, source="platform")
    if other_fee:
        await add_finance("expense", "Biaya Platform", other_fee, f"Biaya {data.platform} {sale['number']}", sale["id"], u, method=data.platform, source="platform")
    if outlet_disc:
        await add_finance("expense", "Diskon Ditanggung Outlet", outlet_disc, f"Diskon {data.platform} {sale['number']}", sale["id"], u, method=data.platform, source="platform")
    await audit("external_sale", sale["id"], "create", u, after={"platform": data.platform, "gross": gross, "net": net, "order_no": data.order_no})
    return sale


# ---------------- Ringkasan keuangan: omzet, alokasi kantong, breakdown metode/channel ----------------
@api.get("/finance/summary")
async def finance_summary(period: str = "month", date: Optional[str] = None, u=Depends(need("finance"))):
    p = rng(period, date)
    fin = await find("finance", {"date": {"$regex": f"^{p}"}}, limit=10000)
    sales = await find("sales", {"date": {"$regex": f"^{p}"}, **VALID_SALE}, limit=10000)
    gross = sum(sale_total(s) for s in sales)
    disc = sum(s.get("discount", 0) for s in sales)
    hpp = sum(s.get("hpp", 0) for s in sales)
    income = sum(f["amount"] for f in fin if f["type"] == "income")
    expense = sum(f["amount"] for f in fin if f["type"] == "expense")
    refund = -sum(f["amount"] for f in fin if f["type"] == "income" and f["amount"] < 0)
    net_omzet = round(gross - refund, 2)
    methods = {}
    for s in sales:
        d = methods.setdefault(s["payment_method"], {"method": s["payment_method"], "count": 0, "total": 0.0, "fee": 0.0})
        d["count"] += 1
        d["total"] += sale_total(s)
    for f in fin:
        if f["type"] == "expense" and f.get("category") in ["Biaya Pembayaran", "Komisi Platform", "Biaya Platform"] and f.get("method") in methods:
            methods[f["method"]]["fee"] += f["amount"]
    for d in methods.values():
        d["fee"] = round(d["fee"], 2)
        d["total"] = round(d["total"], 2)
        d["net"] = round(d["total"] - d["fee"], 2)
    channels = {}
    for s in sales:
        d = channels.setdefault(s.get("channel", "kasir"), {"channel": s.get("channel", "kasir"), "count": 0, "gross": 0.0, "fee": 0.0, "net": 0.0})
        d["count"] += 1
        d["gross"] += sale_total(s)
        d["fee"] += s.get("platform_fee", 0)
        d["net"] += s.get("expected_net", sale_total(s))
    for d in channels.values():
        for k in ["gross", "fee", "net"]:
            d[k] = round(d[k], 2)
    exp_cat = {}
    for f in fin:
        if f["type"] == "expense":
            exp_cat[f["category"]] = round(exp_cat.get(f["category"], 0) + f["amount"], 2)
    pockets = await find("fund_pockets", limit=100)
    cat_pocket = {c["name"]: c.get("pocket_id") for c in await find("expense_categories", limit=200)}
    pocket_actual = {}
    for f in fin:
        if f["type"] == "expense":
            pid = f.get("pocket_id") or cat_pocket.get(f["category"])
            if pid:
                pocket_actual[pid] = round(pocket_actual.get(pid, 0) + f["amount"], 2)
    pocket_rows = []
    for pk in sorted(pockets, key=lambda x: x.get("sort", 0)):
        mth, val = pk.get("alloc_method"), pk.get("alloc_value", 0)
        alloc = round(net_omzet * val / 100, 2) if mth == "percent" else (round(val, 2) if mth == "fixed" else 0.0)
        actual = pocket_actual.get(pk["id"], 0)
        pocket_rows.append({"id": pk["id"], "name": pk["name"], "alloc_method": mth, "alloc_value": val,
                            "allocated": alloc, "actual": round(actual, 2), "diff": round(alloc - actual, 2)})
    return {"period": period, "range": p,
            "omzet": {"gross": round(gross, 2), "discount": round(disc, 2), "refund": round(refund, 2),
                      "net": net_omzet, "hpp": round(hpp, 2), "gross_profit": round(net_omzet - hpp, 2)},
            "cashflow": {"income": round(income, 2), "expense": round(expense, 2), "net": round(income - expense, 2)},
            "by_method": sorted(methods.values(), key=lambda x: -x["total"]),
            "by_channel": sorted(channels.values(), key=lambda x: -x["gross"]),
            "expenses_by_category": [{"category": k, "amount": v} for k, v in sorted(exp_cat.items(), key=lambda x: -x[1])],
            "pockets": pocket_rows}


# ---------------- Config registries (metode bayar, kategori, kantong dana, komponen payroll, platform) ----------------
REGISTRY_FIELDS = {
    "payment_methods": {"name": "", "kind": "other", "fee_percent": 0.0, "fee_fixed": 0.0, "is_active": True, "sort": 0},
    "expense_categories": {"name": "", "description": "", "pocket_id": None, "is_active": True, "sort": 0},
    "fund_pockets": {"name": "", "description": "", "alloc_method": "none", "alloc_value": 0.0, "is_active": True, "sort": 0},
    "payroll_components": {"name": "", "type": "bonus", "description": "", "method": "fixed", "amount": 0.0, "is_active": True, "sort": 0},
    "platforms": {"name": "", "commission_percent": 0.0, "other_fee_percent": 0.0, "fee_fixed": 0.0,
                  "discount_borne_by": "outlet", "payout_schedule": "", "settlement_ref": "", "is_active": True, "sort": 0},
}
REGISTRY_PERM = {"payment_methods": "settings", "expense_categories": "settings", "fund_pockets": "settings",
                 "payroll_components": "payroll", "platforms": "settings"}


def clean_registry(kind, payload):
    out = {}
    for k, dflt in REGISTRY_FIELDS[kind].items():
        v = payload.get(k, dflt)
        if v is None and dflt is not None:
            v = dflt
        if isinstance(dflt, bool):
            v = bool(v)
        elif isinstance(dflt, float):
            try:
                v = float(v)
            except (TypeError, ValueError):
                v = 0.0
        elif isinstance(dflt, int):
            try:
                v = int(v)
            except (TypeError, ValueError):
                v = 0
        out[k] = v
    return out


@api.get("/config/{kind}")
async def config_list(kind: str, active: bool = False, u=Depends(current_user)):
    if kind not in REGISTRY_FIELDS:
        raise HTTPException(404, "Konfigurasi tidak ada")
    items = sorted(await find(kind, limit=500), key=lambda x: (x.get("sort", 0), x.get("name", "")))
    if active:
        items = [x for x in items if x.get("is_active", True)]
    # sembunyikan referensi settlement platform utk non-finance
    if kind == "platforms" and u["role"] not in PERM["settings"]:
        for x in items:
            x.pop("settlement_ref", None)
    return items


@api.post("/config/{kind}")
async def config_create(kind: str, payload: dict, u=Depends(current_user)):
    if kind not in REGISTRY_FIELDS:
        raise HTTPException(404, "Konfigurasi tidak ada")
    if u["role"] not in PERM[REGISTRY_PERM[kind]]:
        raise HTTPException(403, "Akses ditolak untuk jabatan Anda")
    doc = clean_registry(kind, payload)
    if not str(doc["name"]).strip():
        raise HTTPException(400, "Nama wajib diisi")
    created = await insert(kind, doc)
    await audit(kind, created["id"], "create", u, after=doc)
    return created


@api.put("/config/{kind}/{iid}")
async def config_update(kind: str, iid: str, payload: dict, u=Depends(current_user)):
    if kind not in REGISTRY_FIELDS:
        raise HTTPException(404, "Konfigurasi tidak ada")
    if u["role"] not in PERM[REGISTRY_PERM[kind]]:
        raise HTTPException(403, "Akses ditolak untuk jabatan Anda")
    before = await db[kind].find_one({"id": iid, "outlet_id": OUTLET}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Data tidak ditemukan")
    doc = clean_registry(kind, {**before, **payload})
    await db[kind].update_one({"id": iid, "outlet_id": OUTLET}, {"$set": doc})
    await audit(kind, iid, "update", u, before={k: before.get(k) for k in doc}, after=doc)
    return {**before, **doc}


# ---------------- Audit trail viewer ----------------
@api.get("/audit")
async def list_audit(start: Optional[str] = None, end: Optional[str] = None, entity: str = "all",
                     action: str = "all", user_id: str = "all", limit: int = 500, u=Depends(need("audit"))):
    q = {}
    if entity != "all":
        q["entity"] = entity
    if action != "all":
        q["action"] = action
    if user_id != "all":
        q["user_id"] = user_id
    if start or end:
        rng = {}
        if start:
            rng["$gte"] = start
        if end:
            rng["$lte"] = end + "T23:59:59"
        q["created_at"] = rng
    logs = await find("audit_logs", q, limit=min(limit, 2000))
    entities = await db.audit_logs.distinct("entity", {"outlet_id": OUTLET})
    return {"logs": logs, "entities": sorted(entities)}


# ---------------- Seed ----------------
@app.on_event("startup")
async def seed():
    await db.users.create_index("email", unique=True)
    pw = os.environ["DEMO_PASSWORD"]
    owner_email = os.environ["OWNER_EMAIL"].lower()
    if not await db.users.find_one({"email": owner_email}):
        await insert("users", {"email": owner_email, "password_hash": hash_pw(os.environ["OWNER_PASSWORD"]), "name": "Owner ROLLOW",
                               "role": "owner", "position": "Owner", "base_salary": 0, "allowance": 0, "is_active": True})
    demo = [("manager", "Rina Manager", 7000000), ("spv", "Budi Supervisor", 5500000), ("finance", "Sari Finance", 5500000),
            ("marketing", "Dewi Marketing", 5000000), ("crew", "Andi Crew", 3500000)]
    for key, name, sal in demo:
        email = f"{key}@rollow.id"
        if not await db.users.find_one({"email": email}):
            role = "supervisor" if key == "spv" else key
            await insert("users", {"email": email, "password_hash": hash_pw(pw), "name": name, "role": role,
                                   "position": name.split()[1], "base_salary": sal, "allowance": 500000, "is_active": True})
    await db.users.update_many({"pin_hash": {"$exists": False}}, {"$set": {"pin_hash": hash_pw("1234")}})
    await db.users.update_many({"email": "crew@rollow.id", "hourly_rate": {"$exists": False}},
                               {"$set": {"hourly_rate": 15000, "overtime_rate": 20000, "shift_start": "08:00", "base_salary": 0}})
    await db.users.update_many({"shift_start": {"$exists": False}}, {"$set": {"shift_start": "08:00", "hourly_rate": 0, "overtime_rate": 0}})
    # --- config registries (metode bayar, kantong, kategori, komponen gaji, platform) ---
    if await db.payment_methods.count_documents({"outlet_id": OUTLET}) == 0:
        for i, (nm, kd) in enumerate([("Cash", "cash"), ("QRIS", "qris"), ("Debit", "debit"), ("Transfer", "transfer")]):
            await insert("payment_methods", {"name": nm, "kind": kd, "fee_percent": 0.0, "fee_fixed": 0.0, "is_active": True, "sort": i})
    if await db.fund_pockets.count_documents({"outlet_id": OUTLET}) == 0:
        pk = [("HPP / Bahan Baku", "percent", 0), ("Gaji & Payroll", "percent", 0), ("WiFi", "fixed", 0), ("Listrik", "fixed", 0),
              ("Air", "fixed", 0), ("Sewa", "fixed", 0), ("Operasional", "percent", 0), ("Perawatan Alat", "fixed", 0),
              ("Pajak & Admin", "percent", 0), ("Dana Cadangan", "percent", 0)]
        for i, (nm, mth, val) in enumerate(pk):
            await insert("fund_pockets", {"name": nm, "description": "", "alloc_method": mth, "alloc_value": float(val), "is_active": True, "sort": i})
    if await db.expense_categories.count_documents({"outlet_id": OUTLET}) == 0:
        pockets = {p["name"]: p["id"] for p in await find("fund_pockets", limit=100)}
        ec = [("Bahan Baku", "HPP / Bahan Baku"), ("Es Batu", "HPP / Bahan Baku"), ("Sampah / Kebersihan", "Operasional"),
              ("Perlengkapan Operasional", "Operasional"), ("Transportasi", "Operasional"), ("Perbaikan Alat", "Perawatan Alat"),
              ("Gaji", "Gaji & Payroll"), ("WiFi", "WiFi"), ("Listrik", "Listrik"), ("Air", "Air"), ("Sewa", "Sewa"),
              ("Pajak & Admin", "Pajak & Admin"), ("Marketing", "Operasional"), ("Purchasing", "HPP / Bahan Baku"),
              ("Biaya Pembayaran", "Operasional"), ("Lainnya", None)]
        for i, (nm, pk_name) in enumerate(ec):
            await insert("expense_categories", {"name": nm, "description": "", "pocket_id": pockets.get(pk_name), "is_active": True, "sort": i})
    if await db.payroll_components.count_documents({"outlet_id": OUTLET}) == 0:
        comps = [("Bonus Kehadiran", "bonus", "per_day_present", "Nominal per hari hadir"),
                 ("Bonus Kedisiplinan", "bonus", "fixed", "Nominal tetap"),
                 ("Bonus Target", "bonus", "fixed", "Nominal tetap pencapaian target"),
                 ("Bonus Performa", "bonus", "fixed", "Nominal tetap"),
                 ("Potongan Keterlambatan", "deduction", "per_late", "Nominal per keterlambatan"),
                 ("Potongan Pelanggaran SOP", "deduction", "fixed", "Nominal tetap"),
                 ("Potongan Kas/Selisih", "deduction", "manual", "Nominal diisi manual saat input")]
        for i, (nm, tp, mth, desc) in enumerate(comps):
            await insert("payroll_components", {"name": nm, "type": tp, "method": mth, "amount": 0.0, "description": desc, "is_active": True, "sort": i})
    if await db.platforms.count_documents({"outlet_id": OUTLET}) == 0:
        for i, nm in enumerate(["ShopeeFood", "GoFood", "GrabFood"]):
            await insert("platforms", {"name": nm, "commission_percent": 0.0, "other_fee_percent": 0.0, "fee_fixed": 0.0,
                                       "discount_borne_by": "outlet", "payout_schedule": "", "settlement_ref": "", "is_active": True, "sort": i})
    if await db.ingredients.count_documents({}) == 0:
        ings = [("Gula", "gram", 15, 5000, [8000, 1700, 0]), ("Teh Celup", "pcs", 300, 50, [200, 40, 0]),
                ("Kopi Arabika", "gram", 250, 1000, [3000, 800, 0]), ("Susu UHT", "ml", 20, 3000, [10000, 2000, 0]),
                ("Es Batu", "gram", 1, 5000, [0, 10000, 0]), ("Roti Tawar", "pcs", 1500, 20, [40, 0, 15]),
                ("Telur", "pcs", 2200, 30, [60, 0, 20]), ("Kentang Beku", "gram", 40, 2000, [5000, 0, 1500]),
                ("Minyak Goreng", "ml", 18, 2000, [5000, 0, 1000]), ("Coklat Bubuk", "gram", 120, 500, [1500, 300, 0])]
        ids = {}
        sys_user = {"name": "System"}
        for n, unit, cost, mn, stocks in ings:
            d = await insert("ingredients", {"name": n, "unit": unit, "cost_per_unit": cost, "min_stock": mn, "category": "Bahan"})
            ids[n] = d["id"]
            for loc, q in zip(LOCATIONS, stocks):
                if q:
                    await move_stock(d["id"], loc, q, "saldo_awal", "Seed", sys_user)
        menus = [("Es Teh Manis", "Minuman", 8000, "bar", [("Teh Celup", 1), ("Gula", 20), ("Es Batu", 150)]),
                 ("Kopi Susu ROLLOW", "Minuman", 22000, "bar", [("Kopi Arabika", 18), ("Susu UHT", 150), ("Gula", 15), ("Es Batu", 150)]),
                 ("Americano", "Minuman", 18000, "bar", [("Kopi Arabika", 18), ("Es Batu", 150)]),
                 ("Coklat Panas", "Minuman", 20000, "bar", [("Coklat Bubuk", 30), ("Susu UHT", 200), ("Gula", 10)]),
                 ("Roti Bakar Telur", "Makanan", 25000, "kitchen", [("Roti Tawar", 2), ("Telur", 1), ("Minyak Goreng", 10)]),
                 ("French Fries", "Snack", 18000, "kitchen", [("Kentang Beku", 150), ("Minyak Goreng", 50)])]
        for n, cat, price, loc, rec in menus:
            await insert("menus", {"name": n, "category": cat, "price": price, "location": loc, "is_active": True,
                                   "recipe": [{"ingredient_id": ids[i], "qty": q} for i, q in rec]})
        await insert("promos", {"name": "Happy Hour 10%", "kind": "promo", "discount_type": "percent", "value": 10,
                                "start_date": today()[:8] + "01", "end_date": (now() + timedelta(days=60)).strftime("%Y-%m-%d"),
                                "description": "Diskon 10% semua menu", "budget": 0, "is_active": True})


app.include_router(api)
app.add_middleware(CORSMiddleware, allow_credentials=True, allow_origins=["*"], allow_methods=["*"], allow_headers=["*"])
logging.basicConfig(level=logging.INFO)


@app.on_event("shutdown")
async def shutdown_db_client():
    client.close()
